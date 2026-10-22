"""Write docs/speaker_script.pdf: what to say on every slide, and how to explain it.

For each slide of docs/presentation.pptx, in deck order, the script gives
- Say: the words to say, timed at 140 words a minute so the talk fits 15 minutes;
- Formulas and methods: every formula and method on the slide, symbol by symbol;
- If they ask: the questions the slide is likely to raise, with answers.

Like the guide, every number of ours is read from results/ and every number of the
paper's from src/vpin/paper_numbers.py. Slide numbers and titles come from the deck, and
the build fails if a slide has no section, so the script cannot drift from the slides.

Run after make_report.py and make_slides.py (it reads the deck and docs/build/).
Usage: python scripts/make_speaker_script.py
"""
import json
import re
from pathlib import Path

import pandas as pd
from pptx import Presentation
from scipy.stats import norm

from make_guide import CSS, ci, count_tests, pm, render, se, verify
from make_slides import SCORECARD_HIDDEN
from vpin import paper_numbers as PN
from vpin.simulation import BASE, WORLDS

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "speaker_script.pdf"
LABEL = "speaker script"
WORDS_PER_MINUTE = 140
TALK_MINUTES = 15

SCRIPT_CSS = CSS + """
p.f { font-family: Sans; font-size: 10pt; color: #0F2C68; margin: 1pt 0 7pt 16pt; }
p.q { margin: 0 0 7pt 0; }
p.time { font-family: Sans; font-size: 9pt; color: #5a6270; margin: 0 0 6pt 0; }
"""


# ---------- the numbers every section uses ----------

def facts() -> dict:
    """Every number the script quotes, read from results/ (ours) or paper_numbers (the paper's)."""
    k = json.loads((ROOT / "results" / "report" / "key_numbers.json").read_text())
    pred = k["prediction"]
    fc = k["flash_crash"]
    o = k["ive_oos"]
    reg = {(r["market"], r["choice"], r["period"], r["signal"]): r for r in k["regressions"]}
    btc_period = next(r["period"] for r in k["regressions"] if r["market"] == "BTC")
    ledger = pd.DataFrame(k["ledger"])
    claim = lambda start: ledger[ledger["claim"].str.startswith(start)].iloc[0]
    scorecard = pd.read_csv(ROOT / "results" / "report" / "scorecard_ours.csv")
    by_year = pd.read_csv(ROOT / "results" / "extensions" / "ive_by_year.csv")
    oos_vpin = [per["out-of-sample"]["VPIN"]["corr"] for per in o.values()]
    rv_wins = sum(per["out-of-sample"]["RV (trailing volatility)"]["corr"] - 2 * per["out-of-sample"]["RV (trailing volatility)"]["corr_se"]
                  > per["out-of-sample"]["VPIN"]["corr"] for per in o.values())
    kept = [v["placebo_share_of_real_corr"] for v in k["inference"]["placebo"].values()]
    return {
        "k": k, "d": k["data"], "pred": pred, "u250": pred["uncertainty"]["50_250"], "u350": pred["uncertainty"]["10_350"],
        "one_min": fc["BVC, 1-minute bars"], "ten_s": fc["BVC, 10-second bars"], "lr": fc["Lee-Ready, trade by trade"],
        "o": o, "chosen": k["ive_chosen"], "oos_lo": min(oos_vpin), "oos_hi": max(oos_vpin), "rv_wins": rv_wins,
        "bb": k["btc_buckets"]["50_50"], "bse": k["btc_buckets"]["50_50"]["se"], "bars": k["btc_bars"], "tr": k["btc_trades"],
        "b50": k["btc_signals"]["50"],
        "reg_btc": reg[("BTC", "(50, 50)", btc_period, "vpin")],
        "reg_ive": reg[("IVE", "paper (50, 250)", "out-of-sample", "vpin")],
        "sim": lambda world, version, metric: next(r for r in k["simulation"]["summary"]
                                                   if r["world"] == world and r["version"] == version and r["metric"] == metric),
        "dd": k["simulation"]["decisions"], "reps": k["simulation"]["manifest"]["reps"], "sim_days": k["simulation"]["manifest"]["days"],
        "ls": {(r["market"], r["sample_length"], r["signal"]): r for r in k["local_sigma"]["predictive"]},
        "track": k["local_sigma"]["tracking"]["BTC_50_50"],
        "lregs": {(r["market"], r["choice"], r["signal"]): r for r in k["local_sigma"]["regressions"]},
        "inf": k["inference"], "led": k["inference"]["ledger"], "w": k["windows"],
        "n_windows": k["windows"]["IVE"]["windows"] + k["windows"]["BTC"]["windows"],
        "kept_lo": min(kept), "kept_hi": max(kept),
        "flips": sum(v["sign_flips"] for v in k["inference"]["loyo"].values()), "loyo_pairs": len(k["inference"]["loyo"]),
        "placebo_ive": claim("IVE 2011-26, real-time (50, 250): real VPIN predicts better"),
        "placebo_ive_kept": k["inference"]["placebo"]["IVE 2011-26, real-time (50, 250)"]["placebo_share_of_real_corr"],
        "rep_z": claim("corr(ln VPIN, next |r|) > 0")["z"],
        "mcnemar_p": float(norm.sf(claim("VPIN alarm beats the volatility alarm")["z"])),
        "alarms": {(a["market"], a["signal"]): a for a in k["alarms"]},
        "ive_events": [e for e in k["events"] if e["market"] == "IVE"],
        "verdicts": scorecard["Verdict"].value_counts().to_dict(),
        "by_year": by_year, "rv_years": int((by_year["rv_corr"] > by_year["vpin_corr"]).sum()),
        "n_months": len(pd.period_range(k["data"]["period"][0], k["data"]["period"][1], freq="M")),
        "n_tests": count_tests(),
        "pin_base": BASE["alpha"] * BASE["mu"] / (BASE["alpha"] * BASE["mu"] + 2 * BASE["eps"]),
    }


# ---------- page pieces ----------

def deck_slides() -> list[tuple[int, str, bool]]:
    """(number, title, is_backup) for every slide, read from the generated deck."""
    prs = Presentation(ROOT / "docs" / "presentation.pptx")
    out, after_final = [], False
    for i, s in enumerate(prs.slides, 1):
        final = s.slide_layout.name == "Final_1"
        title = "FINAL" if final else (s.shapes.title.text if s.shapes.title is not None else "")
        out.append((i, title, after_final))
        after_final = after_final or final
    return out


def words(html_text: str) -> int:
    return len(re.sub(r"<[^>]+>", " ", html_text).split())


def mmss(seconds: float) -> str:
    minutes, rest = divmod(int(round(seconds)), 60)    # round first, so 59.6 s never prints as 0:60
    return f"{minutes}:{rest:02d}"


def section(number: int, title: str, say: list[str], explain: list[str], qa: list[tuple[str, str]],
            backup: bool = False, when: str = "") -> str:
    """One slide: heading, the words to say, the formulas and methods, the likely questions."""
    shown = "Closing slide" if title == "FINAL" else title
    head = f"<h2>Slide {number} · {shown}</h2>"
    if backup:
        head += f'<p class="time">Backup slide: not presented. Show it when asked {when}</p>'
    else:
        head += f'<p class="time">About {mmss(sum(words(p) for p in say) / WORDS_PER_MINUTE * 60)} at {WORDS_PER_MINUTE} words a minute</p>'
    body = '<p class="label">Say</p><div class="box">' + "".join(f"<p>{p}</p>" for p in say) + "</div>"
    if explain:
        body += '<p class="label">Formulas and methods on this slide</p>' + "".join(explain)
    if qa:
        body += '<p class="label">If they ask</p>' + "".join(f'<p class="q"><b>“{q}”</b> {a}</p>' for q, a in qa)
    return head + body


def item(name: str, text: str) -> str:
    return f"<p><b>{name}.</b> {text}</p>"


def formula(text: str) -> str:
    return f'<p class="f">{text}</p>'


def image(path: str, width: int) -> str:
    return f'<img src="{path}" style="width: {width}%"/>'


# ---------- the slides ----------

def s_title(f, sl):
    say = ["Good afternoon. We are Group 8: Nazarov, Doroshenko, Burlakov, Daunov and Sergeev. Our paper is number 11: "
           "Easley, López de Prado and O'Hara, <i>Flow Toxicity and Liquidity in a High-Frequency World</i>, Review of Financial "
           "Studies, 2012. It introduces VPIN, a real-time measure of toxic order flow. We rebuilt it from scratch on free data, "
           "checked its main claims, and added four extensions."]
    explain = [
        item("VPIN", "Volume-Synchronized Probability of Informed Trading. “Volume-synchronized” because time is measured in "
             "traded volume, not in minutes; “probability of informed trading” because it estimates the share of trading done by "
             "better-informed traders."),
        item("Order-flow toxicity", "The paper's term: order flow is toxic when it adversely selects market makers, who may not "
             "know that they are providing liquidity at a loss."),
    ]
    qa = [("What is new in your work, in one sentence?",
           "We test VPIN where the truth is known, on Bitcoin data with the true buyer and seller of every trade and in a "
           "simulation of the paper's own model, and find that the VPIN the paper computes mostly measures volatility."),
          ("Did you use the authors' code?",
           f"No. All code is ours, written from scratch in Python, with {f['n_tests']} unit tests. The repository is on GitHub "
           f"(backup {sl('Backup: repository')}).")]
    return say, explain, qa


def s_problem(f, sl):
    say = ["Start with who loses when markets break. Market makers, today mostly high-frequency firms, post quotes on both sides "
           "and earn a tiny spread on a huge number of trades. They lose when the trader on the other side knows more: an "
           "informed seller hits their bid just before the price drops, and they are left holding inventory bought too high. "
           "The paper calls such order flow toxic.",
           "When toxicity rises, market makers widen their quotes or stop quoting. With nobody absorbing orders, prices gap. "
           "The authors link this chain to the flash crash of 6 May 2010, when US stock indices fell several percent in minutes "
           "and recovered within the hour. Their question: can we measure toxicity in real time, from trades alone, early "
           "enough to act?"]
    explain = [
        item("Bid–ask spread", "The gap between the best price at which someone will sell (ask) and buy (bid). A market maker "
             "who buys at the bid and sells at the ask earns it on every round trip."),
        item("Adverse selection", "The market maker cannot tell informed from uninformed traders. It loses on average to the "
             "informed and must earn that back from the uninformed through the spread; this is the logic of the classic models "
             "of Glosten and Milgrom (1985) and Kyle (1985). The more informed the flow, the wider the spread must be."),
        item("The four boxes", "The paper's causal chain: one-sided, informed flow → losses on inventory → wider quotes or "
             "withdrawal → a liquidity gap, in which small orders move prices a lot. VPIN is meant to measure the first box."),
    ]
    qa = [("Was the flash crash caused by toxicity?",
           "That is the paper's interpretation. The joint CFTC–SEC staff report (September 2010) describes a large automated "
           "sell programme in E-mini futures meeting thinning liquidity. We do not test the cause; we test whether VPIN warned "
           f"in advance ({sl('Result 1')})."),
          ("Why can't we just observe toxicity?",
           "Trade records say neither who is informed nor, usually, who initiated a trade. That is why the paper needs a proxy, "
           "and why we look for data where the truth is known (Extensions 1 and 3).")]
    return say, explain, qa


def s_pin(f, sl):
    a, mu, eps = BASE["alpha"], BASE["mu"], BASE["eps"]
    say = ["The starting point is PIN, the probability of informed trading, from Easley and co-authors in 1996. News arrives with "
           "probability alpha; informed traders then arrive at rate mu, all on one side; uninformed buyers and sellers arrive "
           "at rate epsilon each. PIN is the informed share of all orders: alpha mu over alpha mu plus two epsilon. It needs "
           "maximum likelihood on daily counts, which is slow and unstable at high frequency.",
           "VPIN replaces the parameters with observable volume. The imbalance between buy and sell volume estimates alpha mu, "
           "because informed traders all trade on one side; total volume estimates the denominator. And time is measured in "
           "volume: trading is cut into buckets of equal volume, one fiftieth of a day. On the right, our data for 6 May 2010: "
           "as trading sped up, the same fifty buckets covered less and less clock time."]
    explain = [
        image("docs/build/eq_pin.png", 30),
        item("Symbols", "α: probability that an information event happens in a period. μ: arrival rate of informed traders when "
             "there is news; they all trade on the side of the news. ε: arrival rate of uninformed buyers, and separately of "
             "uninformed sellers. δ, the probability that news is bad, decides the side but does not enter PIN."),
        item("Why this ratio", "Per period, informed orders arrive only with news, so their expected number is αμ; uninformed "
             "orders number 2ε (buyers plus sellers). PIN = expected informed orders / expected all orders."),
        formula(f"Example (our simulation's baseline): α = {a}, μ = {mu:g}, ε = {eps:g} → PIN = {a * mu:g} / ({a * mu:g} + {2 * eps:g}) = {f['pin_base']:.3f}"),
        item("Maximum likelihood", "PIN is estimated from daily buy and sell counts modelled as a mixture of three Poisson "
             "regimes (no news, good news, bad news). With thousands of trades a day the likelihood overflows numerically and the "
             "optimiser finds local maxima; and it gives one number per day, too slow for intraday use."),
        item("From PIN to VPIN", "In a news period all informed volume is on one side, so the expected |buy − sell| volume is "
             "about μ during news and about 0 without; averaged over periods, about αμ. Total volume is about αμ + 2ε. Their ratio "
             "estimates PIN with nothing fitted: a moment estimator (expectations replaced by sample averages)."),
        item("Volume bucket and volume clock", "A bucket is a slice of trading with exactly V units of volume; V = average daily "
             "volume / 50, so a typical day has 50 buckets and a busy day more. A bar that does not fit into the current bucket "
             "spills into the next. VPIN is updated after every bucket, so it updates faster when trading is busy."),
        item("The figure", "Grey (left axis): how many clock hours the last 50 buckets cover. Blue (right axis): VPIN. On the "
             "afternoon of 6 May 2010 the grey line collapses: the window that normally spans about a day covered a fraction of "
             "an hour."),
    ]
    qa = [("Uninformed buys and sells also differ by chance. Isn't that in |B − S| too?",
           "Yes. In a small bucket chance imbalance adds a positive bias; it shrinks relative to volume as buckets grow. The paper "
           f"works with expected values. In our simulation of exactly this model (Extension 3, {sl('Extension 3')}), VPIN from "
           f"true labels tracks the informed share at {f['dd']['D1_true_vpin_tracks_informed_trading']['M1_T_true']:.2f}, "
           "so the idea works when trades are labelled correctly."),
          ("Why 50 buckets a day?", "It is the paper's baseline. Its Figure 10 scans other choices; we keep the paper's "
           "settings for the reproduction."),
          ("Is VPIN the same as PIN?", "No. PIN is a model parameter estimated by maximum likelihood on daily counts. VPIN is a "
           "running average of volume imbalance in volume time, with no model fitted.")]
    return say, explain, qa


def s_bvc(f, sl):
    z = lambda x: norm.cdf(x)
    say = ["VPIN needs buy and sell volume, but trade data do not say who initiated a trade. Classic rules sign each trade: the "
           "tick rule compares the price with the last different price; Lee and Ready compare it with the quote midpoint. The "
           "paper argues that with order splitting and cancellations these rules break down, and proposes bulk volume "
           "classification.",
           "Trades are grouped into one-minute bars. A share of each bar's volume, the normal CDF of the bar's standardized "
           f"price change, is counted as buys. No change gives fifty-fifty; a rise of one and a half standard deviations gives "
           f"{z(1.5):.0%} buys. It is cheap and robust. But note the last bullet: the imbalance it produces grows with the size "
           "of the price move, by construction. Our extensions come back to exactly this."]
    imb = lambda x: 2 * z(x) - 1
    explain = [
        image("docs/build/eq_bvc.png", 62),
        item("Symbols", "τ: the bucket. i: the one-minute bars inside it. V<sub>i</sub>: volume of bar i (the part that falls in "
             "τ). P<sub>i</sub> − P<sub>i−1</sub>: the bar's price change. σ<sub>ΔP</sub>: the standard deviation of bar price "
             "changes; the paper takes it over the whole sample, our real-time version over the previous 250 days, our fix over "
             "the previous day. Z: the standard normal CDF. Sell volume is the rest of the bucket: V − V<sup>B</sup>."),
        item("Why volatility enters by construction", "Write x = ΔP / σ. A bar's imbalance is |V<sup>S</sup> − V<sup>B</sup>| = "
             "V<sub>i</sub> · |2Z(x) − 1|, which depends only on the size of the move:"),
        formula(f"x = 0 → 0 · x = 1 → {imb(1):.2f} · x = 1.5 → {imb(1.5):.2f} · x = 2 → {imb(2):.2f} (share of the bar's volume counted as one-sided)"),
        item("So", "a turbulent minute looks one-sided whoever traded in it. On a calm market with balanced flow BVC reports "
             "little imbalance; on a volatile market with the same balanced flow it reports a lot."),
        item("Tick rule", "A trade above the last different price is a buy, below it a sell; a trade at the same price keeps the "
             "previous sign."),
        item("Lee–Ready (1991)", "A trade above the midpoint of the bid and ask in force is a buy, below it a sell; at the "
             "midpoint the tick rule decides. Kibot stores the quote at each trade, so no reporting delay is applied."),
        item("Flat bars", f"{f['d']['share_bars_no_price_change']:.0%} of IVE's one-minute bars have no price change, and BVC "
             "splits each of them 50/50, which removes much of the imbalance signal on IVE."),
    ]
    qa = [("Why the normal CDF?", "It is the paper's choice: it maps any price change smoothly to a buy share between 0 and 1. "
           "Any other symmetric CDF keeps the property that matters here: bigger moves mean bigger imbalance."),
          ("Does BVC work?", f"Where the truth is known (Bitcoin, {sl('Extension 1')}) it gets the direction right for "
           f"{f['bars']['direction_agreement']:.1%} of volume, but inflates the imbalance {f['bse']['imbalance_inflation']:.1f} "
           f"times, and its VPIN correlates only {f['bb']['corr_vpin_bvc_true']:.2f} with true-label VPIN."),
          ("What is σ with 10-second bars?", "The standard deviation of 10-second price changes: σ is always computed for the "
           "bar length in use.")]
    return say, explain, qa


def s_data(f, sl):
    d = f["d"]
    say = ["The paper uses E-mini S&amp;P 500 and crude oil futures, 2008 to 2011. Those data are paid. The only free tick history "
           "with quotes we found that covers the flash crash is the IVE ETF from Kibot: every trade with the bid and ask at that "
           f"moment. We reproduce on {d['period'][0]} to {d['period'][1]}: {d['days']} days and {d['ticks']:,} trades, with the "
           "paper's exact settings. For the extensions we add Binance Bitcoin futures, where the true buyer and seller of every "
           "trade are recorded. We remove 42 trades the exchanges cancelled on 6 May, and our ± is a month-block bootstrap.",
           "The bottom table answers last time's question about objectives. VPIN has no training objective. The paper's only "
           "optimisation is the parameter grid search, and it maximises the very statistic it then reports, on the same data. "
           "We select on 2009 to 2011 and evaluate on 2011 to 2026."]
    explain = [
        item("Settings", "V = ADV / 50, where ADV is the average daily volume of the sample; window n = 50 buckets, about one day "
             "of trading; one-minute bars; σ and the CDF from the whole sample. “Identical” on the slide means our reproduction "
             "uses exactly these."),
        item("Whole sample vs real time", "The paper computes σ, V and the CDF over the whole sample, which includes the future. "
             "Our real-time version takes σ and the bucket size for each day from the previous 250 trading days, and ranks each "
             "VPIN value only among values already observed."),
        item("Month-block bootstrap (our ±)", "Split the sample into calendar months; draw the same number of months with "
                                                "replacement; recompute the statistic; repeat 1,000 times (seed 0). ± is the standard "
                                                "deviation of the 1,000 values (a standard error); the 95% interval runs from the "
                                                "2.5th to the 97.5th percentile."),
        item("Why months", "Consecutive VPIN values share 49 of their 50 buckets, and volatility comes in clusters, so buckets are "
             "not independent. Resampling whole months keeps that dependence inside each block."),
        item("The objectives table", "PIN: maximum likelihood of a Poisson mixture. VPIN: a moment estimate, nothing fitted. σ "
             "and V: plain sample statistics. Parameters (buckets per day, window): the grid cell with the highest correlation "
             "(Figure 10) or true-positive rate (Table 5), i.e. a selection objective equal to the reported statistic. Evaluation: "
             "correlation, insurance value, true positives; we add lift and ΔR² over volatility."),
        item("Binance data", "BTCUSDT perpetual futures. One-minute bars (2020–2026) report the taker-buy volume, i.e. the volume "
             "bought by the side that crossed the spread; five days of individual (aggregated) trades carry the aggressor side."),
        item("Cancelled trades", "42 IVE trades at absurd prices (such as $0.08 when the price was about $37), cancelled by the "
             "exchanges after the flash crash under their own rule."),
    ]
    qa = [("What was the objective?", "VPIN has none: it is an average of observed imbalances. The paper's only optimisation "
           "is the grid search over bucket size and window, which maximises the reported statistic on the same data. We select "
           "on 2009–2011 and evaluate on 2011–2026."),
          ("Why IVE and not SPY or the E-mini?", "E-mini tick data are paid. Kibot's free data have a tick history with quotes "
           "only for IVE; SPY minute data are not available to guests."),
          ("Why calendar months as blocks?", f"A month is long compared with the VPIN window, so little dependence crosses block "
           f"boundaries, and the reproduction still has {f['n_months']} months to resample."),
          ("IVE trades 6.5 hours a day. Does the overnight gap matter?", "Without the buckets that span the close the "
           f"correlation is {f['pred']['corr_50_250_intraday']:.2f}, the same as with them.")]
    return say, explain, qa


def s_crash(f, sl):
    one, ten, lr = f["one_min"], f["ten_s"], f["lr"]
    say = ["First result: the flash crash. On the left, the paper's Figure 4(a): on the E-mini, the CDF of VPIN crossed 0.9 at "
           "least two hours before the crash. On the right, the same computation on IVE with three ways of classifying volume.",
           f"With one-minute bulk classification, as in the paper, VPIN is in its top decile from {one['first'][11:16]}, "
           f"{one['lead_hours']:.1f} hours before the fall began at 14:32. That reproduces. With 10-second bars it first crosses "
           f"0.9 at {ten['first'][11:16]}, during the fall. With Lee–Ready, trade by trade, it never does: the maximum is "
           f"{lr['max_cdf']:.2f}. One event allows no ±. The lesson: the warning depends on how volume is classified, and the "
           "paper reports that all three agree."]
    explain = [
        item("CDF(VPIN)", "The share of all VPIN values in the sample that are at or below the current one, from 0 to 1. A value "
             "of 0.9 means VPIN is in its top 10%: the paper's alarm level."),
        item("Lead time", "14:32 New York (the start of the fall in our data; the shaded window runs to 15:08) minus the first "
             "time that day the CDF reached 0.9."),
        item("The three schemes", "BVC on one-minute bars (the paper's baseline); BVC on 10-second bars (same rule, shorter bars, σ "
             "of 10-second changes); Lee–Ready (each trade signed by the quotes, then poured into the same buckets)."),
        item("Why bar length matters (our reading, not tested separately)", "Short bars are more often flat, so split 50/50, and "
             "consecutive short bars often move in opposite directions as trades alternate between bid and ask, so their "
             "imbalances cancel inside a bucket. VPIN then reacts mainly when prices trend, which on 6 May happened during the fall."),
        item("No ±", "The bootstrap resamples months; one afternoon is a single observation, so no standard error is possible."),
    ]
    a = f["alarms"][("IVE", "vpin")]
    qa = [("4.9 hours early: doesn't that show VPIN works?",
           "It shows VPIN was high that morning in one configuration. A warning is only useful if the alarm is not on most of the "
           f"time anyway; in real time on IVE after the paper it is on {a['alarm_share']:.0%} of the time (backup "
           f"{sl('Backup: crashes')})."),
          ("What did Andersen and Bondarenko find?",
           "In “VPIN and the Flash Crash” (2014) they report that on the E-mini VPIN reached its all-time high after the crash, "
           "not before, and that its pre-crash level was elevated but not extraordinary."),
          ("Didn't IVE behave differently from the E-mini that day?",
           "Yes. IVE fell to about half its price for minutes, far more than the index; that is why we also report Table 1 "
           f"without that day (backup {sl('Backup: the volume clock')}).")]
    return say, explain, qa


def s_predict(f, sl):
    u, pred, inf = f["u250"], f["pred"], f["inf"]
    iid = pred["corr_50_250_fisher_ci_iid"]
    say = ["Second, the paper's central predictive claim: high VPIN now means a large move in the next bucket. Figure 10 "
           "correlates log VPIN with the absolute return over the next bucket, for every combination of buckets per day and "
           f"window length. At its preferred setting, 50 buckets a day and a 250-bucket window, the paper reports "
           f"{PN.CORR_50_250:.2f}. On IVE we get {pm(u['corr'])}, with a 95% interval of {ci(u['corr'])}. Same sign, but "
           f"{PN.CORR_50_250:.2f} is far outside.",
           "About the paper's interval: it treats every bucket as independent, but neighbouring VPIN values share 249 of their "
           f"250 buckets. With the paper's formula ours would be [{iid[0]:.2f}, {iid[1]:.2f}]: far too narrow. And it is not a "
           f"question of the data slice: on all {f['n_windows']} two-year windows of IVE and Bitcoin, none reaches "
           f"{PN.CORR_50_250:.2f}, while our design could detect any correlation of {inf['power']['mde_80pct_power']:.2f} or more."]
    explain = [
        formula("corr( ln VPIN<sub>τ−1</sub> , |r<sub>τ</sub>| ),   r<sub>τ</sub> = P<sub>τ</sub> / P<sub>τ−1</sub> − 1"),
        item("Symbols", "P<sub>τ</sub>: the price at the end of bucket τ, so r<sub>τ</sub> is the return over bucket τ. VPIN is "
             "taken at the end of the previous bucket, so it never overlaps the return it predicts. The log is used because VPIN "
             "is skewed."),
        item("The grid", f"Horizontal axis: buckets per day (bucket size = ADV / that number). Vertical axis: the window n in "
             f"buckets. (50, 250) means 50 buckets a day and a window of 250 buckets, about five days. The highest value on our "
             f"grid is {pred['corr_max']:.2f}, at {tuple(pred['corr_argmax'])}."),
        item("The paper's interval (footnote 29)", "Fisher's method: z = artanh(r), interval z ± 1.96 / √(N − 3), transformed back "
             f"with tanh. N is the number of buckets ({PN.CORR_50_250_N:,} in the paper). It assumes N independent pairs; with "
             "overlapping windows the effective sample is far smaller, so the interval is far too narrow."),
        item("Power", f"The smallest correlation the reproduction would detect with 80% probability at 5% two-sided "
             f"significance: MDE = (1.96 + 0.84) × s.e. = 2.8 × {inf['power']['corr_se_50_250']:.3f} = "
             f"{inf['power']['mde_80pct_power']:.2f}."),
        item("Every window", f"The paper's exact in-sample method on every two-year window starting each January and July: "
             f"{f['w']['IVE']['windows']} of IVE, {f['w']['BTC']['windows']} of Bitcoin, all reported (backup "
             f"{sl('Backup: the paper')}). It was run once as an exploration before the pre-registration, and is labelled so. "
             f"The {PN.INSURANCE_50_250:.0%} on the slide is the paper's insurance value, defined on the scorecard slide."),
    ]
    qa = [("Is your 0.08 even significant?", f"Borderline. With the normal approximation z = {f['rep_z']:.2f}; the skewed bootstrap "
           "interval excludes zero. We call it weakly positive."),
          ("Why is yours so much lower?", "Three reasons, which we cannot fully separate without E-mini data. A different market: "
           f"coarse prices, {f['d']['share_bars_no_price_change']:.0%} of bars flat. A calmer period: no 2008 crisis; across a "
           "crisis and calm years VPIN and volatility are both high together, which correlates them even with nothing inside a "
           "regime. And the paper's procedure: look-ahead CDF, grid search, independent-bucket ±."),
          ("Could your code be wrong?", f"The parts that do not depend on prediction reproduce, all {f['n_tests']} unit tests "
           "pass, and our Table 1 formulas reproduce the paper's own Jarque–Bera numbers from its own inputs.")]
    return say, explain, qa


def s_scorecard(f, sl):
    v = {x: c for x, c in f["verdicts"].items() if x not in SCORECARD_HIDDEN}     # only the rows the slide shows
    order = ["reproduced", "partly", "weaker"]
    counts = [f"{v[x]} {x}" for x in order if x in v]
    tally = ", ".join(counts[:-1]) + " and " + counts[-1]
    say = ["This slide collects the main claims, in two separate tables because the data differ. Left: what the paper reports "
           "for the E-mini. Right: what we find on IVE, with ± and a verdict computed only from our data. Reproduced means the "
           "paper's value lies inside our 95% interval and the effect's direction is significant. Weaker means the direction "
           "holds but the paper's value is outside.",
           f"On this slide: {tally}. The pattern is clear. The mechanics reproduce: stability, low VPIN meaning calm, the "
           "one-minute flash-crash warning, and the volume clock in part. The predictive claims are weaker."]
    explain = [
        item("Verdict rule", "Computed on our data only. Reproduced: the paper's value is inside our 95% interval and the direction "
             "is significant (correlation above 0, high VPIN beating low VPIN, lift above 1). Weaker: outside, direction holds. "
             "Partly (the Table 1 row): some Table 1 statistics are significantly better in volume time, not all. The full rule "
             "has two more verdicts, inconclusive (inside but not significant) and not reproduced (otherwise); rows with those "
             "verdicts are not on this slide."),
        item("Not on this slide", f"The flash crash with 10-second bars and with Lee–Ready, both not reproduced, are on "
             f"{sl('Result 1')}. Table 5 at (10, 350), inconclusive, is on backup {sl('Backup: high VPIN')}. The README and the "
             "guide list all ten claims; the row numbers on the slide are renumbered 1, 2, 3, … for the claims shown."),
        item("Stability", "Start the bucket sequence 0, 1, …, 999 one-minute bars later; the standard deviation of VPIN "
             "across these starts, divided by its mean."),
        item("Low VPIN means calm", "P(|r<sub>τ</sub>| &lt; 0.25% | CDF(VPIN<sub>τ−1</sub>) &lt; 0.25)."),
        item("Insurance value", "P(CDF(VPIN<sub>τ−1</sub>) &gt; 0.75 | |r<sub>τ</sub>| &gt; 0.75%): of all big moves, the "
             "share that came right after VPIN was in its top quartile. By chance it would be 25%."),
        item("True positives", "Split the CDF into twenty 5% bands. For every stay of VPIN inside one band, from "
             "entering to leaving it, take the largest absolute return between any two buckets of the stay (the paper's Eq. 10). "
             "A true positive is a stay in the top quartile (bands 15–19) with such a move above 0.75%."),
        item("“P ≥ 0.95” (the Table 1 row)", "Table 1 uses a paired bootstrap: both samplings draw the same months; a statistic "
             "counts as better in volume time if it is lower in at least 95% of replications."),
    ]
    qa = [("Why two tables? One would be easier to read.", "Our instructor's rule: results from different sources are not "
           "comparable cell by cell, here different markets and periods. The row numbers link the two tables."),
          ("Isn't “weaker” a nice word for “failed”?", "It means the direction the paper claims is significant on our data, but "
           "the size is far smaller. “Not reproduced” is kept for cases where the direction itself fails."),
          ("Where are the paper's other claims, like Table 5 at (10, 350)?", f"Table 5 at (10, 350) is inconclusive: ours is "
           f"{pm(f['u350']['tp_top'], True)}, interval {ci(f['u350']['tp_top'], True)}, which contains the paper's "
           f"{PN.TRUE_POSITIVE_10_350:.1%}, but there are only {f['pred']['top_quartile_stays_10_350']} stays in the top quartile "
           f"and high VPIN does not beat low VPIN significantly ({pm(f['u350']['tp_diff'], True)}); backup "
           f"{sl('Backup: high VPIN')}. The flash crash with 10-second bars and Lee–Ready does not reproduce; {sl('Result 1')}.")]
    return say, explain, qa


def s_overfit(f, sl):
    o, chosen = f["o"], f["chosen"]
    best = o["best in-sample true positive"]
    pair = tuple(chosen["best in-sample true positive"])
    by = f["by_year"]
    say = ["The seminar asks how results could be overfit. We see four ways. The CDF that defines an alarm uses the whole sample, "
           "future included. Bucket size and window are picked by scanning grids on the same data, and the best cell is "
           "reported. There is no out-of-sample period, and the ± assumes independent buckets. And the headline evidence is one "
           "flash crash.",
           "Our check removes the first three: a real-time VPIN where σ, bucket size and CDF use only past data, parameters "
           f"frozen on 2009 to 2011, tested on 2011 to 2026. Out of sample every choice lands between {f['oos_lo']:.2f} and "
           f"{f['oos_hi']:.2f}, and trailing volatility, the grey bars, beats VPIN by more than two standard errors for "
           f"{f['rv_wins']} of {len(o)} choices. The best true-positive pair, {pair}, falls from "
           f"{se(best['in-sample']['VPIN'], 'true_positive_top', True)} to {se(best['out-of-sample']['VPIN'], 'true_positive_top', True)}."]
    explain = [
        item("Look-ahead bias", "Using information not available at the time. Ranking a 2010 VPIN value against values from 2011 "
             "tells you, in hindsight, that it was high."),
        item("Selection bias (winner's curse)", "Pick the best of many noisy estimates and its value is inflated by luck; out of "
             "sample it falls back. The chart is the test: in-sample (light) against out-of-sample (dark) bars."),
        item("The four choices", f"The paper's (50, 250); the paper's Table 5 setting (10, 350); our best in-sample correlation "
             f"{tuple(chosen['best in-sample corr'])}; our best in-sample true-positive rate {pair}, both chosen the paper's "
             "way on 2009–2011 and then frozen."),
        item("Real-time pipeline", "For each day σ and the bucket size come from the previous 250 trading days; the CDF ranks each "
             "VPIN value among values already seen. Error bars: ± one month-block standard error."),
        item("Bottom quartile", f"Out of sample, the best pair's true-positive rate when VPIN is in its bottom quartile is "
             f"{se(best['out-of-sample']['VPIN'], 'true_positive_bottom', True)}: low VPIN does as well as high."),
    ]
    qa = [("Does a whole-sample CDF change the correlation?", "No: the correlation uses ln VPIN directly. The CDF matters for "
           "everything defined by thresholds on it: the alarm, the quartiles, the insurance value and the true positives."),
          ("Maybe VPIN simply stopped working after 2011?", f"Possibly. Year by year VPIN's correlation ranges from "
           f"{by['vpin_corr'].min():.2f} to {by['vpin_corr'].max():.2f}, and trailing volatility is higher in {f['rv_years']} of "
           f"{len(by)} years (backup {sl('Backup: real-time')})."),
          ("Isn't your out-of-sample period cherry-picked?", "It is everything after the paper's sample to the end of our data, "
           "August 2011 to September 2026, and the windows analysis reports every two-year window.")]
    return say, explain, qa


def s_e1(f, sl):
    bars, bb, bse, tr = f["bars"], f["bb"], f["bse"], f["tr"]
    say = ["Now our extensions. First: is bulk classification right? The paper could never check, because its data do not record "
           "who initiated each trade. Binance does: every Bitcoin futures trade carries the aggressor's side. So we pour the "
           "bulk-classified and the true volume into the same buckets and compare.",
           f"On one-minute bars BVC gets the direction right for {bars['direction_agreement']:.1%} of volume: better than a coin, "
           f"far from perfect. In buckets it inflates the imbalance {bse['imbalance_inflation']:.1f} times, and the two VPINs "
           f"correlate only {bb['corr_vpin_bvc_true']:.2f}. On five days with every individual trade, the simple tick rule "
           f"classifies {tr['tick_accuracy_min']:.0%} to {tr['tick_accuracy_max']:.0%} of volume correctly. Binance groups the "
           "fills of one order, which helps the tick rule; still, the paper's claim that bulk classification beats trade-by-trade "
           "signing does not hold here."]
    explain = [
        item("Aggressor (taker)", "The side that crossed the spread and so initiated the trade. Binance's one-minute bars report "
             "the taker-buy volume; its individual trades carry a flag “buyer is maker”, which, when true, means the seller was "
             "the aggressor."),
        item("Direction agreement", "For each one-minute bar, does the sign of BVC's imbalance (2Z − 1) match the sign of the true "
             f"imbalance? Volume-weighted share of agreement: {bars['direction_agreement']:.1%} ± "
             f"{bars['direction_agreement_se']:.1%}, between {bars['direction_agreement_min']:.0%} and "
             f"{bars['direction_agreement_max']:.0%} by year."),
        item("Imbalance inflation", f"Mean |sell − buy| per bucket under BVC divided by the same under true labels: "
             f"{bb['mean_abs_imbalance_bvc']:.2f} / {bb['mean_abs_imbalance_true']:.2f} ≈ {bse['imbalance_inflation']:.1f}."),
        item("VPIN correlation", f"Bucket-by-bucket correlation of the two VPIN series: {bb['corr_vpin_bvc_true']:.2f} ± "
             f"{bse['corr_vpin_bvc_true']:.2f} (month blocks)."),
        item("Five trade days", f"Tick-rule accuracy = share of volume whose sign matches the aggressor flag. Correlation of bucket "
             f"imbalance with the truth: tick rule at least {tr['corr_tick_min']:.2f}, BVC on one-minute bars "
             f"{tr['corr_bvc_1min_min']:.2f}–{tr['corr_bvc_1min_max']:.2f}. Ranges are across the five days, so no ±."),
        item("The figure", "Left: BVC VPIN against true-label VPIN, bucket by bucket. Right: both over the last 60 days."),
    ]
    qa = [("Bitcoin is not the E-mini.", "Right. We use it only for questions that need ground truth, is BVC right and is VPIN "
           "more than volatility. Those are properties of the method, not of one market."),
          ("Why does BVC inflate the imbalance?", f"Every price move becomes one-sided volume even when true flow is balanced: a "
           f"bar that rose by one σ is counted as {norm.cdf(1):.0%} buys whoever traded."),
          ("What is an aggregated trade?", "Binance merges the fills of one taker order at the same price into one record. One "
           "large order then shows up as one print, which makes the tick rule's job easier than on equity tapes.")]
    return say, explain, qa


def s_e2(f, sl):
    b, bb, bse = f["b50"], f["bb"], f["bse"]
    rb, ri = f["reg_btc"], f["reg_ive"]
    bvc, true, rv = b["VPIN (bulk classified)"], b["VPIN (true labels)"], b["RV (trailing volatility)"]
    say = ["If bulk classification turns price moves into imbalance, VPIN may be volatility in disguise. Andersen and Bondarenko "
           "argued this in 2014; with true labels we can test it. Same Bitcoin data, same windows, three signals: VPIN from bulk "
           "classification, VPIN from the true labels, and trailing volatility, the average absolute return over the same buckets.",
           f"VPIN from the true labels is negatively related to the next move: {se(true, 'corr')}. Bulk VPIN is positive, "
           f"{se(bvc, 'corr')}, and trailing volatility beats both at {se(rv, 'corr')}. Bulk imbalance moves with the bucket's own "
           f"price change, {bb['corr_abs_imbalance_with_same_bucket_abs_return_bvc']:.2f} against "
           f"{bb['corr_abs_imbalance_with_same_bucket_abs_return_true']:.2f} for the truth: that is where volatility enters. With "
           f"volatility and trading-intensity controls, VPIN adds an R² of {rb['delta_r2']:.4f} on Bitcoin and "
           f"{ri['delta_r2']:.4f} on IVE out of sample."]
    explain = [
        formula("RV<sub>τ</sub> = (1/n) · Σ<sub>j = τ−n+1 … τ</sub> |r<sub>j</sub>|   (trailing volatility, same window and timing as VPIN)"),
        item("Columns", "Correlation of the signal with the next absolute return; P(top quartile | move &gt; 0.75%), the insurance "
             "value; time in top quartile, how often the alarm is on; lift = insurance value ÷ time in top quartile. A signal that "
             "is always on has lift 1."),
        formula("|r<sub>τ</sub>| = a + b<sub>1</sub>·ln VPIN<sub>τ−1</sub> + b<sub>2</sub>·ln RV<sub>τ−1</sub> + b<sub>3</sub>·|r<sub>τ−1</sub>| + b<sub>4</sub>·ln duration<sub>τ−1</sub> + e<sub>τ</sub>   (regressors standardized)"),
        item("Regression", "Duration is the clock time the bucket took to fill (shorter = faster trading), our measure of trading "
             "intensity; for IVE the clock runs only during market hours. The t-statistic of b<sub>1</sub> uses Newey–West "
             "standard errors with n lags, which allow for the overlap of VPIN windows and for heteroskedasticity."),
        item("ΔR²", "R² with VPIN minus R² without it. It cannot be negative, so its bootstrap interval is not a test; the "
             f"t-statistic is. Here t = {rb['t_vpin_hac']:.1f} on Bitcoin (negative) and {ri['t_vpin_hac']:.1f} on IVE."),
    ]
    qa = [("Why is the true-label VPIN negative?", "We have not tested why. One possibility: busy, volatile periods fill buckets "
           "with many independent orders, so true imbalance per unit of volume falls. Andersen and Bondarenko (2014) also found "
           "VPIN built without the mechanical volume link negatively correlated with future volatility."),
          ("On IVE VPIN's t is positive and large. So it does add something?", f"Statistically, on IVE, yes; economically "
           f"ΔR² = {ri['delta_r2']:.4f}, i.e. {ri['delta_r2']:.2%} of the variance. On Bitcoin the sign is the opposite. Small "
           "and of unstable sign."),
          ("Why lift and not the hit rate?", "An alarm that is always on catches every big move. Lift divides by how often the "
           "alarm is on, so it rewards information, not activity.")]
    return say, explain, qa


def s_e3(f, sl):
    sim, dd = f["sim"], f["dd"]
    v, t = WORLDS["V"], WORLDS["T"]
    holds = {"D1": dd["D1_true_vpin_tracks_informed_trading"]["holds"],
             "D2": dd["D2_false_alarms_from_volatility"]["paper"]["confounds_volatility"],
             "D3": dd["D3_predictive_test_uninformative"]["holds"], "D4": dd["D4_local_sigma_is_an_improvement"]["holds"]}
    held, failed = [r for r, ok in holds.items() if ok], [r for r, ok in holds.items() if not ok]
    rules = (", ".join(held[:-1]) + " and " + held[-1] + " hold" if len(held) > 1 else f"{held[0]} holds") + \
        ("; " + " and ".join(failed) + " does not, and we report it" if failed else "")
    say = ["Extension 3 closes a gap in the paper's own validation. Its Monte Carlo draws buys and sells straight from the PIN "
           "model: every trade correctly labelled, no prices. But its empirical VPIN uses bulk classification, which needs "
           "prices. So the method it actually uses is never checked. We simulated the paper's own model and added a price: "
           "order flow moves it, public news adds noise.",
           "Four worlds. In T, informed trading rises in some blocks of days; in V only volatility and uninformed volume rise; TV "
           f"has both, 0 neither. {f['reps']} replications each, with hypotheses and decision rules written down before the first "
           f"run. The key number, world V: the paper's VPIN tracks the true informed share at "
           f"{sim('V', 'paper', 'M1_tracking')['mean']:.2f}; with true labels {sim('V', 'true', 'M1_tracking')['mean']:.2f}; with "
           f"our fix {sim('V', 'local', 'M1_tracking')['mean']:.2f}. {rules}. When "
           "volatility changes, the paper's VPIN stops measuring informed trading."]
    explain = [
        item("Model", f"One step is one minute of a market that trades around the clock; {f['sim_days']} days per replication; "
             f"information periods of 60 minutes. News with probability α = {BASE['alpha']}, good or bad equally likely; informed "
             f"traders arrive at Poisson rate μ = {BASE['mu']:g} per minute on the news side; uninformed buys and sells at rate "
             f"ε = {BASE['eps']:g} each; every trade is one unit."),
        formula("Δp<sub>t</sub> = λ · (B<sub>t</sub> − S<sub>t</sub>) + σ<sub>pub</sub> · z<sub>t</sub>,   z<sub>t</sub> ~ N(0, 1)"),
        item("Price", f"λ = {BASE['lam'] * 1e4:g} bp per unit of net order flow: permanent price impact, so one-sided informed flow "
             f"moves the price (the paper's toxicity-induced volatility). σ<sub>pub</sub> = {BASE['sigma'] * 1e4:g} bp a minute: "
             "public news nobody trades on."),
        item("Worlds", f"T: α = {t.alpha_high} instead of {BASE['alpha']} in alternating {t.t_block_days}-day blocks. V: "
             f"σ<sub>pub</sub> = {v.sigma_high * 1e4:g} bp and ε = {v.eps_high:g} in {v.v_block_days}-day blocks, informed "
             "trading unchanged. TV: both, on different schedules. 0: the baseline throughout."),
        item("Three VPINs", "Paper: BVC with σ over the whole replication. True: the true buy volume. Local (our fix): BVC with σ "
             "from the previous day only. True informed share: informed volume ÷ all volume over the same n buckets."),
        item("Metrics", "M1 tracking: correlation of VPIN with the true informed share. M2 false alarms (world V): share of buckets "
             "with CDF ≥ 0.9 in high blocks minus low blocks. M3 detection (world T): AUC, the probability that a random "
             "high-α bucket scores higher than a random low-α bucket (0.5 = chance). M4: the paper's predictive correlation."),
        item("Monte Carlo ±", "Standard deviation across replications ÷ √R."),
        item("Decision rules", "D1: true-label VPIN tracks informed trading if M1 &gt; 0.5 in world T. D2: a VPIN confounds "
             "volatility if M2 &gt; 0 by more than 3 Monte Carlo s.e. D3: the predictive test is fooled if M4 &gt; 0 by more than "
             "3 s.e. in world V. D4: the fix improves if its M2 is under half the paper's and its AUC at most 0.05 lower."),
        item("Why true-label VPIN falls in V", f"Uninformed volume doubles (ε {BASE['eps']:g} → {v.eps_high:g}), so the informed "
             "share falls; correct labels follow it. The same negative sign as on real Bitcoin data."),
        item("Why D3 failed", "Two effects cancel in the paper's predictive statistic: volatility inflates the bulk imbalance, extra "
             "volume dilutes it."),
    ]
    d2 = dd["D2_false_alarms_from_volatility"]["paper"]
    qa = [("Was the simulation built to give your answer?", "Its design and decision rules were written down before the first "
           "run (docs/preregistration.md; its hash is in every run manifest), and one pre-registered prediction, D3, failed. "
           "The model is the paper's own PIN model plus the simplest possible price."),
          ("Why these parameter values?", f"They were fixed in the pre-registration before any run; the baseline PIN is "
           f"{f['pin_base']:.3f}. World V triples public volatility and doubles uninformed volume, as on turbulent days. We did "
           "not search over them."),
          ("All three VPINs detect toxicity (AUC near 1). So the paper's VPIN works?", "It detects real toxicity in world T. The "
           f"problem is that it also fires when only volatility changes: false alarms {d2['M2_V']:+.3f} ± {d2['mc_se']:.3f} in "
           "world V. An alarm that fires for both cannot tell them apart.")]
    return say, explain, qa


def s_e4(f, sl):
    ls, tr, lregs = f["ls"], f["track"], f["lregs"]
    btc_s, btc_f = ls[("BTC", 50, "VPIN, standard BVC")], ls[("BTC", 50, "VPIN, local-sigma BVC (fix)")]
    ive_s, ive_f = ls[("IVE 2011-26", 250, "VPIN, standard BVC")], ls[("IVE 2011-26", 250, "VPIN, local-sigma BVC (fix)")]
    li, lb = lregs[("IVE 2011-26", "(50, 250)", "vpin_local")], lregs[("BTC", "(50, 50)", "vpin_local")]
    say = ["Because D4 held, the pre-registration required us to try the fix on real data: bulk classification with σ from the "
           "previous day only, so each price change is judged against current volatility, not against the whole sample.",
           f"Two results. On Bitcoin, where the truth is known, the fixed VPIN tracks true-label VPIN better: "
           f"{tr['corr_local_true']['value']:.2f} against {tr['corr_standard_true']['value']:.2f}, a difference of "
           f"{tr['difference']['value']:+.2f} ± {tr['difference']['se']:.2f}. And once the volatility bias is gone, VPIN's "
           f"predictive correlation collapses: on Bitcoin from {btc_s['corr']:.3f} to {btc_f['corr']:.3f}, on IVE from "
           f"{ive_s['corr']:.3f} to {ive_f['corr']:.3f}. What remains beyond volatility is small: a t of {li['t_vpin_hac']:.1f} on "
           "IVE, a negative coefficient on Bitcoin. The paper's predictive power came from the bias, not from toxicity."]
    explain = [
        formula("buy share = Z( ΔP / σ<sub>previous day</sub> )"),
        item("Local σ", "The root mean square of the previous day's one-minute log price changes: 1,440 bars for Bitcoin, which "
             "trades around the clock, and 390 for IVE, one 6.5-hour session. Before the first full day the buy share is 0.5. "
             "Standard real-time BVC uses σ from the previous 250 days."),
        item("Why it removes the bias", "On a turbulent day a large move is ordinary relative to that day's σ, so it is not counted "
             "as strongly one-sided. With a long-run σ every turbulent minute looks extreme and becomes fake imbalance."),
        item("Tracking", "Bucket-by-bucket correlation with true-label VPIN on Bitcoin; the difference is a paired month-block "
             "bootstrap on the same buckets."),
        item("The table", "Real-time pipeline, 50 buckets a day. Correlation with the next absolute return and lift, ± month-block "
             "standard errors; orange marks the fix."),
        item("Beyond volatility", f"The same regression as Extension 2: IVE t = {li['t_vpin_hac']:.1f}, ΔR² = "
             f"{li['delta_r2']:.4f}; Bitcoin t = {lb['t_vpin_hac']:.1f}."),
    ]
    qa = [("Isn't this circular? You removed volatility, so of course it predicts volatility less.",
           "The point is which version is closer to toxicity. Where the truth is known, the fixed VPIN tracks true-label VPIN "
           "better. So the extra predictive power of the paper's version is volatility, not toxicity."),
          ("Why the previous day and not the previous hour?", "One day was fixed in the pre-registration before we saw any result; "
           "we did not try other windows."),
          ("So is the fixed VPIN useful?", "It is a better measure of toxicity where we can check. It also shows how little "
           "toxicity tells us about next-bucket volatility in these data.")]
    return say, explain, qa


def s_adds(f, sl):
    led, inf, p = f["led"], f["inf"], f["placebo_ive"]
    say = ["What does this add to the paper? One: we close its validation gap; in its own model, VPIN as the paper computes it "
           "stops measuring informed trading when volatility changes. Two: we locate VPIN's predictive power. With true labels it "
           f"disappears, and with the direction of every imbalance shuffled at random VPIN keeps {f['kept_lo']:.0%} to "
           f"{f['kept_hi']:.0%} of its correlation: it is the size of price moves, not one-sided trading.",
           "Three: a one-line fix that tracks true toxicity better where it can be checked. Four: an honest evaluation, out of "
           f"sample, every window, pre-registered, ± everywhere, and a ledger of all {led['tests']} tests: {led['supported_bh']} "
           f"supported after Benjamini–Hochberg, {led['supported_bonferroni']} after Bonferroni, {led['significant_wrong_side']} "
           "significant on the wrong side, counted as failures. Ground truth and the volatility result confirm Andersen and "
           "Bondarenko on new data; the simulation and the fix are new."]
    explain = [
        item("Placebo", "For every one-minute bar keep the size of its BVC imbalance, |2f − 1| × volume, but draw its sign at random "
             "(20 draws, seeds 0–19); recompute VPIN; compare its correlation with the next absolute return. Kept = placebo "
             "correlation ÷ real correlation. Pre-registered rule: keeping 80% or more means direction carries little of VPIN's "
             "predictive power."),
        item("Leave one year out", f"Recompute correlation and lift dropping each year in turn; report the minimum, the maximum and "
             f"sign flips: {f['flips']} flips across {f['loyo_pairs']} signal–sample pairs."),
        item("Power", f"MDE = (1.96 + 0.84) × s.e. = {inf['power']['mde_80pct_power']:.2f}: the paper's 0.40 would not be missed."),
        item("Ledger", "One row per test behind a claim: estimate, standard error, z or t, two-sided p-value. A claim is supported "
             "only if significant and on the claimed side."),
        formula(f"Bonferroni: reject if p ≤ 0.05 / m, m = {led['tests']}. Controls the chance of even one false claim."),
        formula("Benjamini–Hochberg: sort p<sub>(1)</sub> ≤ … ≤ p<sub>(m)</sub>; take the largest i with p<sub>(i)</sub> ≤ (i / m) × 0.05; "
                "reject the first i. Controls the expected share of false claims among those supported."),
        item("Also flagged", "Harvey–Liu–Zhu's |t| &gt; 3, a stricter bar proposed for finance research."),
    ]
    qa = [("What is new compared with Andersen and Bondarenko?", "They used perfect trade classification on E-mini data and argued "
           "that VPIN's link to volatility is mechanical; our Bitcoin results confirm this on new data. What we add: a "
           "known-answer simulation of the paper's own model, a fix tested against ground truth, the direction placebo, and "
           "pre-registered inference."),
          ("On IVE 2011–26 the real VPIN beats the placebo significantly. Doesn't that contradict you?", f"It is significant but "
           f"tiny: {p['estimate']:.4f} ± {p['se']:.4f} in correlation, and the placebo still keeps {f['placebo_ive_kept']:.1%}. "
           "Direction carries a little information there, almost none of VPIN's predictive power."),
          ("Why report both Benjamini–Hochberg and Bonferroni?", "Bonferroni guards against any false claim and is conservative; "
           "Benjamini–Hochberg controls the share of false claims and is the usual choice for many tests. Showing both lets the "
           "reader pick.")]
    return say, explain, qa


def s_critique(f, sl):
    say = ["Our critique, limited to what we tested. Bulk classification builds volatility into VPIN; with true labels the "
           "predictive link disappears. A whole-sample CDF and grid-searched parameters bring look-ahead and selection, with no "
           "out-of-sample test. The main evidence is one event, and on IVE even that depends on bar length and classifier. There "
           "is no baseline: the paper never compares VPIN with plain volatility, which wins in all our tests. Its ± treats "
           "overlapping buckets as independent. And the authors disclose a patent and a financial interest in VPIN, one more "
           "reason for independent replication.",
           "With three more months we would: rerun everything on E-mini futures with exchange aggressor flags, the paper's own "
           "market with ground truth; test VPIN as a filter for a market-making strategy after costs; correct the grid search "
           "for multiple testing; and combine VPIN with a HAR volatility forecast to test its value out of sample."]
    explain = [
        item("Patent", "The paper's title page discloses that the authors hold a patent on VPIN and a financial interest in it. It "
             "does not make the results wrong; it is a reason to replicate independently."),
        item("Exchange aggressor flags", "CME's MDP 3.0 market data report the aggressor side of each trade, so E-mini VPIN could be "
             "computed from true labels: the paper's market with ground truth."),
        item("Market-making filter", "Simulate a market maker that widens or pulls quotes when VPIN is high; measure adverse-selection "
             "losses (how prices move against its fills) and profit after costs, with and without the filter."),
        item("White's reality check (2000)", "Tests whether the best of many tried specifications really beats a benchmark, by "
             "bootstrapping the maximum statistic over all of them, so the search itself is accounted for."),
        item("Deflated statistics", "For example the deflated Sharpe ratio (Bailey and López de Prado, 2014), which discounts the "
             "best result by the number of trials. López de Prado is one of the paper's authors."),
        formula("HAR (Corsi, 2009):  RV<sub>t+1</sub> = c + β<sub>d</sub> RV<sub>t</sub> + β<sub>w</sub> RV<sub>t</sub><sup>week</sup> "
                "+ β<sub>m</sub> RV<sub>t</sub><sup>month</sup> + e<sub>t+1</sub>"),
        item("HAR in words", "Tomorrow's volatility from today's, last week's average and last month's average. A strong, standard "
             "benchmark; VPIN is useful only if it adds to it out of sample."),
    ]
    qa = [("Isn't this unfair to a first proposal?", "We credit what holds: the mechanics reproduce. The critique is about the "
           "evidence for the predictive claims, which is where the paper's procedure is weakest."),
          ("Which next step first?", "E-mini with aggressor flags. It is the paper's own market with ground truth, so it would "
           "separate the market from the method."),
          ("What would change your conclusion?", "A clear out-of-sample advantage of VPIN over trailing volatility on the E-mini "
           "with true aggressor flags.")]
    return say, explain, qa


def s_final(f, sl):
    say = ["Thank you. In one sentence: the mechanics of VPIN reproduce, but the VPIN the paper computes mostly measures "
           "volatility, and plain volatility predicts better. We are happy to take your questions."]
    explain = [item("Backup slides, by question", "Where is VPIN's formula and how is it read? "
                    f"{sl('Backup: the VPIN metric')}. Table 1 and stability? {sl('Backup: the volume clock')}. Tables 4 and 5? "
                    f"{sl('Backup: high VPIN')}. Every two-year window? {sl('Backup: the paper')}. Real-time year by year? "
                    f"{sl('Backup: real-time')}. Crashes after the paper? {sl('Backup: crashes')}. Simulation metrics? "
                    f"{sl('Backup: the known-answer')}. Code and how to run it? {sl('Backup: repository')}."),
               item("If you do not know", "Say what we did not test and what would answer the question. It is better than a guess, "
                    "and the instructor values honesty about limits.")]
    return say, explain, []


# ---------- backups ----------

def b_metric(f, sl):
    t2 = f["k"]["table2"]
    say = ["VPIN after bucket τ is the total absolute imbalance of the last n buckets divided by their total volume, n × V. "
           "Because its level depends on the market and the settings, the paper reads it through its CDF: where today's value "
           "sits among all values."]
    explain = [image("docs/build/eq_vpin.png", 42),
               item("Symbols", "V<sup>S</sup><sub>j</sub>, V<sup>B</sup><sub>j</sub>: sell and buy volume of bucket j; n: the window "
                    "(50 buckets ≈ one day); V: bucket size; nV: the volume of the window."),
               item("Levels", f"Mean VPIN on IVE is {t2['ours']['Average']:.2f}; the paper's Table 2 reports "
                    f"{t2['paper']['Average']:.2f} for the E-mini. Coarse IVE prices make many bars flat, split 50/50, so imbalances "
                    "are smaller; the CDF makes the two comparable."),
               item("The figure", "VPIN over the reproduction period; the flash-crash afternoon stands out.")]
    return say, explain, [("Why not read VPIN's level directly?", "Its level depends on bar size, classification and the market; "
                           "the CDF turns it into a rank, which is comparable across settings.")]


def b_clock(f, sl):
    k = f["k"]
    sig = k["table1_significant"]
    say = ["Table 1 asks whether price changes sampled in volume time look more like independent normal draws than changes sampled "
           "every so many minutes. Ours drops 6 May 2010 and overnight changes. Volume time is significantly better on "
           f"{sig['50']} of 4 statistics at 50 draws a day and {sig['100']} of 4 at 100: partly reproduced. Stability "
           f"reproduces: {k['stability']['ours']:.1%} of the mean against the paper's {k['stability']['paper']:.1%}."]
    explain = [
        item("Sampling", "Clock time: the last price in each of 50 (or 100) equal intervals of the day. Volume time: the price each "
             "time cumulative volume crosses a multiple of ADV / 50 (or ADV / 100). Same number of draws, so the statistics are "
             "comparable."),
        item("Statistics (on standardized price changes)", "Kurt: raw kurtosis, 3 for a normal distribution. J-B*: the Jarque–Bera "
             "statistic divided by T, (s² + (k − 3)² / 4) / 6, with skewness s and kurtosis k. L-B*: the Ljung–Box Q statistic at "
             "10 lags (autocorrelation). White*: R² of x<sub>t</sub>² on all products x<sub>t−i</sub>x<sub>t−j</sub>, "
             "1 ≤ i ≤ j ≤ 10 (heteroskedasticity). Max: the largest standardized change. Lower means closer to independent normal."),
        item("P(volume lower)", "Paired bootstrap: both samplings draw the same months; the share of 1,000 replications in which "
             "volume time has the lower value. Green cells: at least 0.95."),
        item("Why drop 6 May 2010", "One IVE price change that day is about 100 standard deviations and dominates every statistic."),
        item("Stability (Figure 5)", "Restart the buckets 0, 1, …, 999 one-minute bars later; the standard deviation of VPIN across "
             "these starts, divided by its mean."),
    ]
    return say, explain, [("Your Table 1 formula: how do you know it is the paper's?", "Using the paper's own skewness and "
                           "kurtosis, our J-B* formula reproduces the paper's own J-B* numbers; a unit test checks it.")]


def b_tables45(f, sl):
    u, u3 = f["u250"], f["u350"]
    say = [f"Two separate heat maps: the paper's Table 4a and ours. Low VPIN means calm in both: after bottom-quartile VPIN the next "
           f"move is below 0.25% {u['bottom_small']['value']:.1%} of the time on IVE. The warning claims are weaker: the insurance "
           f"value is {pm(u['insurance'], True)}, still about {u['lift']['value']:.1f} times chance."]
    explain = [item("Reading the heat map", "Each row is a 5% band of CDF(VPIN<sub>τ−1</sub>); each cell is the share of next-bucket "
                    "absolute returns in a 0.25% bin; rows sum to 1."),
               item("Insurance and true positives", "As defined on the scorecard slide: P(top quartile | move &gt; 0.75%), and the "
                    "share of top-quartile stays with a move above 0.75%."),
               item("Table 5 at (10, 350)", f"Ours {pm(u3['tp_top'], True)}, interval {ci(u3['tp_top'], True)}: wide, because there "
                    f"are only {f['pred']['top_quartile_stays_10_350']} stays.")]
    return say, explain, []


def b_windows(f, sl):
    w = f["w"]
    say = [f"The paper's exact in-sample method on every two-year window starting each January and July: {w['IVE']['windows']} "
           f"windows of IVE, {w['BTC']['windows']} of Bitcoin. On IVE the correlation ranges from {w['IVE']['corr_50_250'][0]:.2f} "
           f"to {w['IVE']['corr_50_250'][1]:.2f}, on Bitcoin from {w['BTC']['corr_50_250'][0]:.2f} to "
           f"{w['BTC']['corr_50_250'][1]:.2f}. No window reaches the paper's numbers."]
    explain = [item("In-sample method", "σ, bucket size and CDF from the window itself, as the paper does over its whole sample."),
               item("Bitcoin caveat", "A 0.75% move is routine on Bitcoin, so the true-positive rate is as high or higher when VPIN is "
                    "low; that statistic says nothing there."),
               item("Status", "Run once as an exploration before the pre-registration; every window is reported.")]
    return say, explain, []


def b_by_year(f, sl):
    by = f["by_year"]
    say = [f"Real-time VPIN at the paper's (50, 250), year by year from 2011. VPIN's correlation with the next move ranges from "
           f"{by['vpin_corr'].min():.2f} to {by['vpin_corr'].max():.2f}; trailing volatility is higher in {f['rv_years']} of "
           f"{len(by)} years. The share of time the alarm is on drifts a lot, because a CDF from the past does not fit the future."]
    explain = [item("Alarm rate", "The share of buckets with real-time CDF above 0.75; with a stable distribution it would stay "
                    "near 25%.")]
    return say, explain, []


def b_crashes(f, sl):
    a, r = f["alarms"][("IVE", "vpin")], f["alarms"][("IVE", "rv")]
    days = ", ".join(e["day"] for e in f["ive_events"])
    pv = "below 0.001" if a["p_value_binomial"] < 0.001 else f"{a['p_value_binomial']:.3f}"
    say = [f"Was the real-time alarm on when each day's largest one-hour fall began? On IVE ({days}) the VPIN alarm was on at "
           f"{a['hits']} of {a['falls']}, while it is on {a['alarm_share']:.0%} of the time; trailing volatility at {r['hits']}. "
           f"This is the result most favourable to VPIN, but on the same falls VPIN does not beat volatility significantly "
           f"(McNemar p = {f['mcnemar_p']:.2f}), and several falls sit inside weeks of turmoil."]
    explain = [item("Alarm", "Real-time CDF ≥ 0.9 at the start of the day's largest one-hour fall; 50 buckets a day, n = 50."),
               item("Binomial p", f"The chance of at least that many hits if the alarm were on at random with its overall rate, "
                    f"falls treated as independent: {pv} for VPIN on IVE. A rough guide, because falls cluster."),
               item("McNemar test (exact)", "Compares two alarms on the same events using only the events where they disagree "
                    "(one on, the other off); with few such events it has little power.")]
    return say, explain, []


def b_sim_metrics(f, sl):
    say = ["Every simulation metric with its Monte Carlo error bar: false alarms from volatility in world V, detection of toxic "
           "episodes in world T, and the paper's predictive statistic where informed trading never changes."]
    explain = [item("Error bars", "± 1.96 Monte Carlo standard errors (standard deviation across replications ÷ √R)."),
               item("Reading", "Blue: the paper's method; green: true labels; orange: our fix. Left: blue above zero is D2. Middle: "
                    "all near 1, detection works. Right: blue slightly below zero, so D3 did not hold.")]
    return say, explain, []


def b_repo(f, sl):
    say = [f"All code is ours, written from scratch in Python. One command per stage; {f['n_tests']} unit tests; every number in the "
           "README, slides, guide and this script is read from results/. Raw data are downloaded by scripts, not committed."]
    explain = [item("Order", "pytest; reproduce.py; extensions.py; windows.py; simulation.py; local_sigma.py; inference.py; "
                    "make_report.py; make_slides.py; make_guide.py; make_speaker_script.py."),
               item("Where to look", "DEVIATIONS.md lists every deviation from the paper with its reason; docs/preregistration.md "
                    "the pre-registered tests and change log.")]
    return say, explain, []


# Title prefix of each slide -> its section; deck order comes from the .pptx itself.
SECTIONS = [
    ("11. FLOW", s_title), ("Why toxic", s_problem), ("Method: from PIN", s_pin), ("Method: bulk", s_bvc),
    ("Data, and", s_data), ("Result 1", s_crash), ("Result 2", s_predict), ("Reproduction scorecard", s_scorecard),
    ("How could", s_overfit), ("Extension 1", s_e1), ("Extension 2", s_e2), ("Extension 3", s_e3), ("Extension 4", s_e4),
    ("What our work adds", s_adds), ("Critique and next", s_critique), ("FINAL", s_final),
    ("Backup: the VPIN metric", b_metric), ("Backup: the volume clock", b_clock), ("Backup: high VPIN", b_tables45),
    ("Backup: the paper", b_windows), ("Backup: real-time", b_by_year), ("Backup: crashes", b_crashes),
    ("Backup: the known-answer", b_sim_metrics), ("Backup: repository", b_repo),
]

WHEN = {"Backup: the VPIN metric": "how VPIN is computed or read.", "Backup: the volume clock": "about Table 1 or stability.",
        "Backup: high VPIN": "about Tables 4 and 5.", "Backup: the paper": "whether another period would give the paper's numbers.",
        "Backup: real-time": "whether VPIN worked after the paper.", "Backup: crashes": "about later crashes.",
        "Backup: the known-answer": "for the full simulation results.", "Backup: repository": "about the code."}


def build_html() -> tuple[str, float]:
    f = facts()
    slides = deck_slides()
    titles = {title: number for number, title, _ in slides}
    sl = lambda prefix: f"slide {next(n for t, n in titles.items() if t.startswith(prefix))}"
    parts, talk_seconds, plan = [], 0.0, []
    for number, title, backup in slides:
        prefix, build = next(((p, b) for p, b in SECTIONS if title.startswith(p)), (None, None))
        if build is None:
            raise SystemExit(f"slide {number} ({title}) has no section in the speaker script")
        say, explain, qa = build(f, sl)
        if not backup:
            seconds = sum(words(p) for p in say) / WORDS_PER_MINUTE * 60
            talk_seconds += seconds
            plan.append(f"<tr><td>{number}</td><td>{'Closing slide' if title == 'FINAL' else title}</td>"
                        f"<td>{mmss(seconds)}</td><td>{mmss(talk_seconds)}</td></tr>")
        parts.append(section(number, title, say, explain, qa, backup, WHEN.get(prefix, "")))
    n_main = sum(not b for _, _, b in slides)
    front = f"""
<h1>Speaker script: what to say on every slide</h1>
<p class="sub">Flow toxicity and VPIN · Research Seminar · Group 8 · {n_main} slides in the talk and {len(slides) - n_main} backups</p>
<div class="box">
<p><b>How to use it.</b> Each slide has three parts. <b>Say</b> is the text to speak, written for {WORDS_PER_MINUTE} words a minute
so that the whole talk takes about {mmss(talk_seconds)} of the {TALK_MINUTES} minutes. <b>Formulas and methods</b> explains every
formula and method on the slide, symbol by symbol, so you can answer a question about any of them. <b>If they ask</b> lists the
questions the slide is likely to raise, with answers. Numbers are read from the project's results, and slide numbers from the
deck, so they match the slides exactly.</p>
</div>
<p class="label">Talk plan</p>
<table><tr><th>Slide</th><th>Title</th><th>Time</th><th>Cumulative</th></tr>{''.join(plan)}</table>
"""
    return front + "".join("<!--pagebreak-->" + p for p in parts), talk_seconds


def main():
    html_text, talk_seconds = build_html()
    if talk_seconds > TALK_MINUTES * 60:
        raise SystemExit(f"the script takes {mmss(talk_seconds)}, more than {TALK_MINUTES} minutes: shorten a Say text")
    pages = render(html_text, OUT, label=LABEL, title="Flow toxicity and VPIN: speaker script", css=SCRIPT_CSS)
    verify(html_text, OUT, label=LABEL)
    print(f"saved {OUT} ({pages} pages, talk {mmss(talk_seconds)}), all text blocks present")


if __name__ == "__main__":
    main()
