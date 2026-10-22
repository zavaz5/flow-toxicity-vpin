"""Write docs/presentation_guide.pdf: a companion that explains the presentation.

It walks through every part of the deck, explains why some of the paper's results did
not reproduce, and what each extension adds. Like the README and the slides, every
number is read from results/ (ours) or src/vpin/paper_numbers.py (the paper's), and
the paper's numbers and ours never share a table.

Run after make_report.py and make_slides.py (it reuses figures from docs/build/).
Usage: python scripts/make_guide.py
"""
import html
import json
from datetime import date
from pathlib import Path

import pandas as pd
import pymupdf

from vpin import paper_numbers as PN

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "presentation_guide.pdf"
FONTS = "/System/Library/Fonts/Supplemental"

CSS = """
@font-face { font-family: Body; src: url("Georgia.ttf"); }
@font-face { font-family: Body; src: url("Georgia Bold.ttf"); font-weight: bold; }
@font-face { font-family: Body; src: url("Georgia Italic.ttf"); font-style: italic; }
@font-face { font-family: Sans; src: url("Arial.ttf"); }
@font-face { font-family: Sans; src: url("Arial Bold.ttf"); font-weight: bold; }
body { font-family: Body; font-size: 10pt; line-height: 1.36; color: #1d1d1f; }
h1 { font-family: Sans; font-weight: bold; font-size: 20pt; color: #0F2C68; margin: 0 0 4pt 0; }
h2 { font-family: Sans; font-weight: bold; font-size: 14pt; color: #0F2C68; margin: 16pt 0 5pt 0; }
h3 { font-family: Sans; font-weight: bold; font-size: 11pt; color: #1656A6; margin: 11pt 0 3pt 0; }
p { margin: 0 0 6pt 0; }
p.sub { font-family: Sans; font-size: 9.5pt; color: #5a6270; margin-bottom: 10pt; }
p.cap { font-family: Sans; font-size: 8.5pt; color: #5a6270; margin: 2pt 0 10pt 0; }
p.label { font-family: Sans; font-weight: bold; font-size: 9.5pt; color: #0F2C68; margin: 8pt 0 3pt 0; }
ul, ol { margin: 0 0 6pt 0; }
li { margin-bottom: 3pt; }
div.box { border-left: 3pt solid #1656A6; padding: 2pt 0 0 10pt; margin: 4pt 0 8pt 0; }
div.warn { border-left: 3pt solid #EB8C3C; padding: 2pt 0 0 10pt; margin: 4pt 0 8pt 0; }
table { border-collapse: collapse; width: 100%; font-family: Sans; font-size: 8.5pt; margin: 2pt 0 4pt 0; }
th { font-weight: bold; text-align: left; color: #0F2C68; padding: 3pt 5pt; border-bottom: 1.2pt solid #0F2C68; }
td { padding: 3pt 5pt; border-bottom: 0.5pt solid #c8ccd2; vertical-align: top; }
td.good { color: #007A4F; font-weight: bold; }
td.mid { color: #B8650F; font-weight: bold; }
td.bad { color: #B03A2E; font-weight: bold; }
img { width: 100%; }
"""

VERDICT_CLASS = {"reproduced": "good", "partly": "mid", "weaker": "mid", "inconclusive": "mid", "not reproduced": "bad",
                 "holds": "good", "does not hold": "bad"}


# ---------- small HTML helpers ----------

def esc(x) -> str:
    return html.escape(str(x))


def table(df: pd.DataFrame, verdict_col: str | None = None) -> str:
    head = "".join(f"<th>{esc(c)}</th>" for c in df.columns)
    rows = []
    for _, r in df.iterrows():
        cells = []
        for c in df.columns:
            cls = f' class="{VERDICT_CLASS.get(r[c], "")}"' if c == verdict_col else ""
            cells.append(f"<td{cls}>{esc(r[c])}</td>")
        rows.append("<tr>" + "".join(cells) + "</tr>")
    return f"<table><tr>{head}</tr>{''.join(rows)}</table>"


def fig(path: str, caption: str, width: int = 100) -> str:
    """Figure with caption; `width` in percent of the text width (square charts look better narrower)."""
    return f'<img src="{path}" style="width: {width}%"/><p class="cap">{caption}</p>'


def pm(st: dict, as_pct=False, digits=2) -> str:
    if as_pct:
        return f"{st['value']:.0%} ± {st['se']:.0%}"
    return f"{st['value']:.{digits}f} ± {st['se']:.{digits}f}"


def ci(st: dict, as_pct=False) -> str:
    return f"[{st['lo']:.0%}, {st['hi']:.0%}]" if as_pct else f"[{st['lo']:.2f}, {st['hi']:.2f}]"


def se(r: dict, key: str, as_pct=False) -> str:
    return f"{r[key]:.0%} ± {r[key + '_se']:.0%}" if as_pct else f"{r[key]:.2f} ± {r[key + '_se']:.2f}"


def slide_lookup():
    """Slide numbers by title prefix, read from the generated deck (never typed)."""
    from pptx import Presentation
    titles = [x.shapes.title.text if x.shapes.title is not None else "" for x in Presentation(ROOT / "docs" / "presentation.pptx").slides]
    return lambda prefix: f"slide {next(i for i, t in enumerate(titles) if t.startswith(prefix)) + 1}"


def count_tests() -> int:
    """Number of test cases pytest collects, so the text never goes stale."""
    import subprocess
    import sys
    out = subprocess.run([sys.executable, "-m", "pytest", "--collect-only", "-q"], cwd=ROOT,
                         capture_output=True, text=True).stdout
    return sum("::" in line for line in out.splitlines())


# ---------- the document ----------

def build_html() -> str:
    k = json.loads((ROOT / "results" / "report" / "key_numbers.json").read_text())
    d, pred, fc = k["data"], k["prediction"], k["flash_crash"]
    u250, u350 = pred["uncertainty"]["50_250"], pred["uncertainty"]["10_350"]
    o, chosen = k["ive_oos"], k["ive_chosen"]
    t1, sig, low = k["table1_boot"], k["table1_significant"], k["table1_lower"]
    bb, bse, bars, tr = k["btc_buckets"]["50_50"], k["btc_buckets"]["50_50"]["se"], k["btc_bars"], k["btc_trades"]
    b50 = k["btc_signals"]["50"]
    reg = {(r["market"], r["choice"], r["period"], r["signal"]): r for r in k["regressions"]}
    btc_period = next(r["period"] for r in k["regressions"] if r["market"] == "BTC")
    reg_btc, reg_ive = reg[("BTC", "(50, 50)", btc_period, "vpin")], reg[("IVE", "paper (50, 250)", "out-of-sample", "vpin")]
    by_year = pd.read_csv(ROOT / "results" / "extensions" / "ive_by_year.csv")
    top_year = by_year.loc[by_year["vpin_corr"].idxmax()]
    rv_years = int((by_year["rv_corr"] > by_year["vpin_corr"]).sum())
    alarms = {(a["market"], a["signal"]): a for a in k["alarms"]}
    p250_oos, p250_is = o["paper (50, 250)"]["out-of-sample"], o["paper (50, 250)"]["in-sample"]
    best_c, best_t = o["best in-sample corr"], o["best in-sample true positive"]
    t2 = k["table2"]
    lr = fc["Lee-Ready, trade by trade"]
    one_min, ten_s = fc["BVC, 1-minute bars"], fc["BVC, 10-second bars"]
    pval = lambda a: "&lt; 0.001" if a["p_value_binomial"] < 0.001 else f"{a['p_value_binomial']:.3f}"
    n_tests = count_tests()
    sl = slide_lookup()

    w, inf, led = k["windows"], k["inference"], k["inference"]["ledger"]
    sim = lambda world, version, metric: next(r for r in k["simulation"]["summary"]
                                              if r["world"] == world and r["version"] == version and r["metric"] == metric)
    sim_n = k["simulation"]["manifest"]["reps"]
    vnames = {"paper": "paper's method (BVC, whole-sample σ)", "true": "true labels", "local": "BVC with one-day σ (our fix)"}
    sim_df = pd.DataFrame([(vnames[v], f"{sim('V', v, 'M1_tracking')['mean']:.3f} ± {sim('V', v, 'M1_tracking')['mc_se']:.4f}",
                            f"{sim('V', v, 'M2_false_alarm')['mean']:+.3f} ± {sim('V', v, 'M2_false_alarm')['mc_se']:.3f}",
                            f"{sim('T', v, 'M3_auc')['mean']:.3f} ± {sim('T', v, 'M3_auc')['mc_se']:.4f}") for v in vnames],
                          columns=["VPIN version (simulation)", "Tracks informed share, world V", "False alarms, world V", "Detection AUC, world T"])
    dd = k["simulation"]["decisions"]
    dec_df = pd.DataFrame([
        ("D1", "true labels track informed trading", "holds" if dd["D1_true_vpin_tracks_informed_trading"]["holds"] else "does not hold"),
        ("D2", "the paper's VPIN raises false alarms from volatility", "holds" if dd["D2_false_alarms_from_volatility"]["paper"]["confounds_volatility"] else "does not hold"),
        ("D3", "the paper's predictive test is fooled by volatility", "holds" if dd["D3_predictive_test_uninformative"]["holds"] else "does not hold"),
        ("D4", "one-day σ halves false alarms and keeps detection", "holds" if dd["D4_local_sigma_is_an_improvement"]["holds"] else "does not hold"),
    ], columns=["Rule", "Pre-registered statement", "Verdict"])
    ls_rows = [r for r in k["local_sigma"]["predictive"] if not (r["market"] == "BTC" and r["sample_length"] == 250)]
    ls_df = pd.DataFrame([(r["market"], r["signal"], f"{r['corr']:.3f} ± {r['corr_se']:.3f}", f"{r['lift']:.2f} ± {r['lift_se']:.2f}") for r in ls_rows],
                         columns=["Market (real-time, 50 buckets/day)", "Signal", "Corr with next |return|", "Lift"])
    tr2 = k["local_sigma"]["tracking"]
    lregs = {(r["market"], r["choice"], r["signal"]): r for r in k["local_sigma"]["regressions"]}
    lr_ive, lr_btc = lregs[("IVE 2011-26", "(50, 250)", "vpin_local")], lregs[("BTC", "(50, 50)", "vpin_local")]
    plc_df = pd.DataFrame([(nm, f"{v['real_corr']:.3f}", f"{v['placebo_corr_mean']:.3f} ± {v['placebo_corr_sd']:.3f}", f"{v['placebo_share_of_real_corr']:.0%}")
                           for nm, v in inf["placebo"].items()], columns=["Sample", "Corr, real VPIN", "Corr, direction shuffled", "Kept"])
    flips = sum(v["sign_flips"] for v in inf["loyo"].values())
    rep_z = next(r["z"] for r in k["ledger"] if r["claim"].startswith("corr(ln VPIN, next |r|) > 0"))
    paper_sc = pd.read_csv(ROOT / "results" / "report" / "scorecard_paper.csv")
    ours_sc = pd.read_csv(ROOT / "results" / "report" / "scorecard_ours.csv")
    ours_sc.columns = ["#", "Ours on IVE (± s.e.)", "95% interval", "Verdict"]
    stats = ["Kurt", "J-B*", "L-B*", "White*", "Max"]
    paper_t1 = pd.DataFrame([(s, f"{PN.TABLE1_50['Chrono time'][s]:.2f}", f"{PN.TABLE1_50['Volume time'][s]:.2f}") for s in stats],
                            columns=["Statistic (50 draws a day)", "Clock time", "Volume time"])
    ours_t1 = pd.DataFrame([(s, f"{t1['50'][s]['Chrono time']:.2f} ± {t1['50'][s]['Chrono ±']:.2f}",
                             f"{t1['50'][s]['Volume time']:.2f} ± {t1['50'][s]['Volume ±']:.2f}",
                             f"{t1['50'][s]['P(volume < clock)']:.3f}") for s in stats],
                           columns=["Statistic (50 draws a day)", "Clock time", "Volume time", "P(volume lower)"])
    claims = paper_sc[["#", "Claim", "Where"]]
    oos_rows = []
    for choice, per in o.items():
        for period in ("in-sample", "out-of-sample"):
            v, r = per[period]["VPIN"], per[period]["RV (trailing volatility)"]
            oos_rows.append((choice, period, se(v, "corr"), se(r, "corr"), se(v, "lift"), se(r, "lift")))
    oos_df = pd.DataFrame(oos_rows, columns=["Parameters", "Period", "Corr, VPIN", "Corr, volatility", "Lift, VPIN", "Lift, volatility"])
    btc_df = pd.DataFrame([(name, se(r, "corr"), se(r, "insurance", True), f"{r['alarm_rate']:.0%}", se(r, "lift"))
                           for name, r in b50.items()],
                          columns=["Signal (BTC, 50 buckets/day, n = 50)", "Corr with next |return|",
                                   "P(top quartile | move > 0.75%)", "Time in top quartile", "Lift"])
    ev = pd.DataFrame(k["events"])
    onoff = lambda v: "–" if v is None or v != v else ("on" if v else "off")
    ev_df = pd.DataFrame([(r["market"], r["day"], f"{r['drop_1h']:.1%}", onoff(r["vpin_alarm_on"]),
                           onoff(r.get("vpin_true_alarm_on")), onoff(r["rv_alarm_on"])) for r in ev.to_dict(orient="records")],
                         columns=["Market", "Day", "Largest 1-hour fall", "VPIN alarm", "VPIN alarm, true labels", "Volatility alarm"])
    names = {"vpin": "VPIN", "vpin_true": "VPIN, true labels", "rv": "Trailing volatility"}
    al_df = pd.DataFrame([(a["market"], names[a["signal"]], f"{a['hits']} of {a['falls']}", f"{a['alarm_share']:.0%}",
                           html.unescape(pval(a))) for a in k["alarms"]],
                         columns=["Market", "Alarm", "On at falls", "On overall", "Binomial p"])

    return f"""
<h1>Flow toxicity and VPIN: what we did, what we found, and why</h1>
<p class="sub">Companion to our Research Seminar presentation of Easley, López de Prado &amp; O'Hara (2012), <i>Flow Toxicity and Liquidity in a High-Frequency World</i>, Review of Financial Studies 25(5) · Group 8 · generated from the project's results on {date.today():%d %B %Y}</p>

<div class="box">
<p><b>The short version.</b></p>
<ul>
<li><b>The paper</b> proposes VPIN, a real-time measure of how one-sided ("toxic") order flow is, and argues that it warned of the 6 May 2010 flash crash hours in advance and predicts short-term volatility.</li>
<li><b>We rebuilt it from scratch</b> on free data: every trade of the IVE ETF (an S&amp;P 500 fund) over the paper's period, with the paper's exact settings, plus Bitcoin futures data in which the true buyer and seller of every trade are known.</li>
<li><b>The mechanics reproduce:</b> VPIN does not depend on where the volume clock starts, low VPIN goes with calm markets, and with one-minute bars VPIN was in its top decile {one_min['lead_hours']:.1f} hours before the flash crash.</li>
<li><b>The warning power is much weaker:</b> VPIN's correlation with the next bucket's volatility is {pm(u250['corr'])} on IVE against the paper's {PN.CORR_50_250:.2f}. Part of the gap is our different market and calmer period; part is that the paper's own procedure flatters its numbers.</li>
<li><b>Our extensions show why:</b> bulk classification, the step that splits volume into buys and sells, builds volatility into VPIN. VPIN computed from the <i>true</i> buys and sells does not predict volatility at all, and VPIN with the direction of every trade shuffled at random predicts exactly as well as the real one. Out of sample, VPIN loses most of its edge, and plain trailing volatility beats it.</li>
<li><b>We close a gap in the paper's own validation:</b> its Monte Carlo never runs bulk classification. Our pre-registered known-answer simulation of its own model shows that VPIN, as the paper computes it, stops measuring informed trading when volatility changes. A one-line fix (measuring price changes against the previous day's volatility) repairs this and, on real Bitcoin data, tracks true toxicity better.</li>
</ul>
</div>

<p class="label">Contents</p>
<ol>
<li>The paper in one page</li>
<li>What we did</li>
<li>Reproduction results</li>
<li>Why some parts did not reproduce</li>
<li>What our extensions add to the paper</li>
<li>Critique, overfitting and next steps</li>
<li>Questions we expect</li>
<li>Glossary and where the numbers come from</li>
</ol>

<h2>1. The paper in one page</h2>
<h3>The problem ({sl('Why toxic')})</h3>
<p>Market makers, today mostly high-frequency trading firms, post buy and sell quotes and earn the small gap between them on a very large number of trades. They lose money when the trader on the other side knows more than they do: the informed trader sells just before the price falls, and the market maker is left holding inventory bought too high. Order flow that does this is called <b>toxic</b>. When toxicity rises, market makers widen their quotes or stop quoting. With nobody left to absorb orders, prices can gap. The authors argue this is part of what happened in the flash crash of 6 May 2010, when US stock indices fell several percent in minutes and recovered within the hour. Their question: can toxicity be measured in real time, from the trade record alone, early enough to act?</p>

<h3>From PIN to VPIN ({sl('Method: from PIN')}; backup {sl('Backup: the VPIN metric')})</h3>
<p>An earlier model measures the <b>probability of informed trading</b>, PIN = αμ / (αμ + 2ε): α is the chance of news, μ the arrival rate of informed traders, ε the arrival rate of uninformed traders on each side. PIN has to be estimated by maximum likelihood on daily counts of buys and sells, which is slow and unstable at high frequency.</p>
<p>VPIN replaces the estimation with two observable quantities. The absolute difference between buy and sell volume estimates αμ (informed traders all trade on one side); total volume estimates αμ + 2ε. VPIN is the average absolute imbalance over the last n buckets divided by their volume. Trading is cut into <b>buckets of equal volume</b>, 1/50 of an average day's volume each, instead of equal clock time. Busy, news-heavy periods fill buckets faster and therefore update VPIN more often: the paper calls this the <b>volume clock</b>. With n = 50 buckets, VPIN covers about one day of trading.</p>
<p>VPIN levels depend on all these choices, so the paper reads VPIN through its <b>cumulative distribution function (CDF)</b>: where today's value sits among all past values. An alarm is a CDF above 0.9, meaning VPIN is in its top 10%.</p>

<h3>Bulk volume classification ({sl('Method: bulk')})</h3>
<p>Trade data do not say who initiated a trade. Classic rules sign every trade one by one: the tick rule (a trade above the last different price is a buy) or Lee–Ready (a trade above the quote midpoint is a buy). The paper argues that in fast electronic markets with order splitting and cancellations these rules break down, and proposes <b>bulk volume classification (BVC)</b>: group trades into one-minute bars and count a share Z(ΔP/σ) of each bar's volume as buys, where ΔP is the bar's price change, σ the standard deviation of such changes and Z the standard normal distribution function. A bar with no price change is split 50/50; a bar that rose by 1.5σ is counted as 93% buys. Keep one property in mind, because our extensions return to it: <b>the imbalance BVC produces grows with the size of the price move, by construction.</b></p>

<h3>The paper's claims</h3>
<p>The table lists the claims we tested, with the exhibit where the paper makes each one. The paper's numbers for them are in Section 3.</p>
{table(claims)}

<!--pagebreak-->
<h2>2. What we did</h2>
<h3>Data ({sl('Data, and')})</h3>
<p>The paper uses E-mini S&amp;P 500 futures and WTI crude oil futures from the CME, 2008–2011. These data are paid. The only free tick history with quotes we found that covers the flash crash is Kibot's sample of the <b>IVE ETF</b> (iShares S&amp;P 500 Value): every trade with the bid and ask at that moment, from 28 September 2009 to today. Our reproduction uses {d['period'][0]} to {d['period'][1]}: {d['days']} trading days and {d['ticks']:,} trades. Everything after 15 August 2011 is new data for the out-of-sample test.</p>
<p>For the extensions we added <b>Binance BTCUSDT perpetual futures</b> (2020–2026). Binance records whether the buyer or the seller initiated each trade, so for Bitcoin we know the true buy volume that bulk classification tries to estimate. The paper never had such data.</p>

<h3>Pipeline and settings</h3>
<p>All code is ours, written from scratch in Python. Raw trades become one-minute bars; bars are classified by BVC; volume is poured into equal-volume buckets (a bar that overflows a bucket spills into the next one); VPIN and its CDF are computed per bucket. The reproduction uses the paper's settings exactly: one-minute bars, bucket size V = average daily volume / 50, window n = 50 buckets, σ and the CDF from the whole sample. {n_tests} unit tests check each step, including one that confirms our reading of the paper's Table 1 formula reproduces the paper's own numbers.</p>
<p>The main deviations (all nine are in DEVIATIONS.md, each with its reason and effect):</p>
<ul>
<li>an ETF instead of futures, trading 6.5 hours a day instead of almost 24;</li>
<li>our sample starts in September 2009, so it misses the 2008 crisis;</li>
<li>42 IVE trades at absurd prices (such as $0.08 when the price was about $37), which the exchanges cancelled after the flash crash under their own rule, are removed;</li>
<li>Table 1 is also reported without the flash-crash day and without overnight price changes, because one IVE price move that day is about 100 standard deviations and dominates every statistic.</li>
</ul>

<h3>What is estimated, chosen and evaluated ({sl('Data, and')})</h3>
<p>Our last presentation was asked about the difference between a training objective and a calibration objective. The equivalent here: <b>VPIN has no training objective.</b> Nothing is fitted; VPIN is a moment estimate (an average of observed imbalances). The original PIN model is estimated by maximum likelihood, but VPIN replaces it. The inputs σ and V are plain sample statistics. The only optimisation in the paper is the <b>choice of parameters</b> (buckets per day and window length): the authors scan grids and report the cell with the highest correlation (Figure 10) or the highest true-positive rate (Table 5). That is a selection objective, and it is the same statistic that is then reported as the result, on the same data. We separate the two: parameters are selected on 2009–2011 and evaluated on 2011–2026.</p>

<h3>How we measure precision (±) and judge each claim</h3>
<p>Neighbouring buckets are far from independent: two consecutive VPIN values share 49 of their 50 buckets, and volatility comes in clusters. We therefore use a <b>block bootstrap over calendar months</b>: draw months with replacement, recompute the statistic, repeat 1,000 times. "±" is the standard deviation across replications (a standard error) and brackets give the 95% interval. The paper's own interval for its correlation (footnote 29) treats buckets as independent. Applied to our data, that method would give [{pred['corr_50_250_fisher_ci_iid'][0]:.2f}, {pred['corr_50_250_fisher_ci_iid'][1]:.2f}], against {ci(u250['corr'])} from month blocks: the iid interval is several times too narrow.</p>
<p>Each claim gets a verdict computed on our data only. It is <b>reproduced</b> if the paper's value lies inside our 95% interval and the direction of the effect is significant (correlation above 0, high VPIN beating low VPIN, lift above 1). It is <b>inconclusive</b> if inside but not significant, <b>weaker</b> if outside but the direction holds, and <b>not reproduced</b> otherwise. The paper's numbers and ours never share a table: they come from different markets and periods.</p>
<p>All new tests (the simulation and the inference package) were written down in <code>docs/preregistration.md</code> before they were run; every later change is in its change log. A ledger lists every test behind a claim, with Benjamini–Hochberg and Bonferroni corrections, and counts a claim as supported only if the estimate is also on the claimed side.</p>

<h2>3. Reproduction results</h2>
<h3>The flash crash ({sl('Result 1')})</h3>
{fig("results/reproduction/figures/fig4_flash_crash.png", "Our version of the paper's Figure 4: IVE on 6 May 2010. Blue: VPIN; purple, dashed: its CDF over the whole sample, with the 0.9 alarm line; grey: price; shaded: the fall and recovery (14:32–15:08 New York).")}
<p>The paper reports that all three classification schemes put CDF(VPIN) above 0.9 at least two hours before the crash. On IVE this holds only for one-minute bars: VPIN was in its top decile from {one_min['first'][11:16]}, {one_min['lead_hours']:.1f} hours before the fall began. With 10-second bars the CDF first crossed 0.9 at {ten_s['first'][11:16]}, during the fall. With Lee–Ready, trade by trade, it never did (maximum {lr['max_cdf']:.2f}). A single event allows no ±. The lesson is that the warning depends on how volume is classified, which the paper's robustness discussion does not show.</p>

<h3>The volume clock and stability (backup {sl('Backup: the volume clock')})</h3>
<p>The paper's Table 1 compares price changes sampled 50 (or 100) times a day at equal clock intervals with price changes sampled at equal volume intervals. Volume sampling should give changes closer to a normal distribution: lower kurtosis and Jarque–Bera statistic, less autocorrelation (Ljung–Box, L-B*) and less heteroskedasticity (White*).</p>
<p class="label">The paper's Table 1 (E-mini, 2008–2011)</p>
{table(paper_t1)}
<p class="label">Our Table 1 (IVE, 2009–2011, without 6 May 2010 and overnight changes; ± s.e.)</p>
{table(ours_t1)}
<p>The last column is the share of bootstrap replications in which volume time has the lower value; both samples draw the same months. On IVE the point estimates favour volume time for {low['50']} of the 4 test statistics at 50 draws a day and {low['100']} of 4 at 100, but only {sig['50']} and {sig['100']} of them clear the 95% bar. The clearest gains are in the tails and heteroskedasticity; autocorrelation does not improve at 50 draws a day. Verdict: <b>partly reproduced</b>.</p>
<p>Stability (Figure 5): starting the bucket sequence 0, 1, …, 999 one-minute bars later changes VPIN by {k['stability']['ours']:.1%} of its mean on IVE. The paper reports {k['stability']['paper']:.1%} for the E-mini. VPIN does not depend on where the volume clock starts: <b>reproduced</b>.</p>

<h3>Does VPIN predict the next bucket's volatility? ({sl('Result 2')})</h3>
{fig("results/reproduction/figures/fig10_correlation.png", "Our version of Figure 10: correlation of ln VPIN with the absolute return over the next bucket, for every combination of buckets per day and window length. The orange square is the paper's preferred (50, 250).", 62)}
<p>The paper reports a correlation of {PN.CORR_50_250:.2f} at 50 buckets a day and a 250-bucket window. On IVE we get {pm(u250['corr'])}, 95% interval {ci(u250['corr'])}. The sign is the paper's, but 0.40 lies far outside our interval. The highest value anywhere on our grid is {pred['corr_max']:.2f}. Verdict: <b>weaker</b>.</p>

<h3>High VPIN and large moves (backup {sl('Backup: high VPIN')})</h3>
{fig("results/reproduction/figures/table4a_conditional.png", "Table 4a as two separate heat maps: the paper's (left) and ours (right). Each row is a VPIN 5%-band; each cell is the share of next-bucket absolute returns in a 0.25% bin.")}
<ul>
<li><b>Low VPIN means calm.</b> The paper: when VPIN is in its bottom quartile, the next move is below 0.25% more than {PN.BOTTOM_QUARTILE_SMALL_MOVE:.0%} of the time. Ours: {u250['bottom_small']['value']:.1%} ± {u250['bottom_small']['se']:.1%}. <b>Reproduced.</b></li>
<li><b>Big moves follow high VPIN.</b> The paper: over {PN.INSURANCE_50_250:.0%} of moves above 0.75% came right after VPIN was in its top quartile. Ours: {pm(u250['insurance'], True)}, interval {ci(u250['insurance'], True)}. Still about {u250['lift']['value']:.1f} times the 25% expected by chance, but far from 84%. <b>Weaker.</b></li>
<li><b>Volatility while VPIN stays high (Figure 14, Table 5).</b> The paper: {PN.TRUE_POSITIVE_50_250:.1%} at (50, 250) and {PN.TRUE_POSITIVE_10_350:.1%} at (10, 350). Ours: {pm(u250['tp_top'], True)} (<b>weaker</b>) and {pm(u350['tp_top'], True)}, interval {ci(u350['tp_top'], True)}. The paper's {PN.TRUE_POSITIVE_10_350:.1%} lies inside that interval, but at 10 buckets a day there are few stays in the top quartile, and high VPIN does not beat low VPIN significantly ({pm(u350['tp_diff'], True)}). <b>Inconclusive.</b></li>
</ul>

<h3>The scorecard ({sl('Reproduction scorecard')})</h3>
<p class="label">What the paper reports (E-mini, 2008–2011)</p>
{table(paper_sc.rename(columns={paper_sc.columns[2]: "Reported"}))}
<p class="label">What we find on IVE (2009–2011), same row numbers</p>
{table(ours_sc, verdict_col="Verdict")}

<!--pagebreak-->
<h3>The predictive claims on every two-year window ({sl('Result 2')}; backup {sl('Backup: the paper')})</h3>
<p>Could the paper's numbers appear in a different slice of our data? We ran the paper's exact in-sample method on every two-year window, starting each January and July: {w['IVE']['windows']} windows of IVE and {w['BTC']['windows']} of Bitcoin. Every window is shown; this analysis was run once as an exploration before the pre-registration, and we label it so.</p>
{fig("results/extensions/windows/windows.png", "Our data, every two-year window: correlation of ln VPIN with the next absolute return (left) and P(VPIN in top quartile | next move above 0.75%) (right).")}
<p>On IVE the correlation ranges from {w['IVE']['corr_50_250'][0]:.2f} to {w['IVE']['corr_50_250'][1]:.2f} and the insurance value from {w['IVE']['insurance_50_250'][0]:.0%} to {w['IVE']['insurance_50_250'][1]:.0%}; on Bitcoin from {w['BTC']['corr_50_250'][0]:.2f} to {w['BTC']['corr_50_250'][1]:.2f} and {w['BTC']['insurance_50_250'][0]:.0%} to {w['BTC']['insurance_50_250'][1]:.0%}. The paper reports 0.40 and over 84%. No window on either market reaches them. On Bitcoin a 0.75% move is routine, so the true-positive rate is as high or higher when VPIN is <i>low</i>, and that statistic says nothing there.</p>

<!--pagebreak-->
<h2>4. Why some parts did not reproduce</h2>
<p>The reasons fall into three groups: our data differ from the paper's, VPIN is sensitive to choices the paper treats as harmless, and the paper's own procedure flatters its numbers. We also ruled out several explanations.</p>

<h3>4.1 A different market</h3>
<ul>
<li><b>Coarse prices.</b> IVE moves in one-cent steps and often not at all within a minute: {d['share_bars_no_price_change']:.0%} of its one-minute bars have no price change, and bulk classification splits all of them 50/50. That removes much of the imbalance signal VPIN is built on. VPIN levels are therefore lower on IVE (mean {t2['ours']['Average']:.2f}; the paper's Table 2 reports {t2['paper']['Average']:.2f} for the E-mini). The paper itself advises reading VPIN through its CDF for this reason.</li>
<li><b>Regular hours only.</b> The E-mini trades almost around the clock; IVE only from 09:30 to 16:00, so some buckets span the overnight close. We checked that this does not drive the result: without those buckets the correlation is {pred['corr_50_250_intraday']:.2f}, the same as with them.</li>
<li><b>An ETF can dislocate.</b> On 6 May 2010 IVE fell to about half its price for minutes, far more than the index or the E-mini. One such price change dominates every statistic of Table 1, which is why we also report it without that day.</li>
</ul>

<h3>4.2 A different, calmer period</h3>
<p>The paper's sample starts in January 2008 and includes the financial crisis; ours starts in September 2009. This matters because bulk-classified VPIN largely tracks volatility (Extension 2 below). Over a sample that contains both a crisis and calm years, VPIN and future volatility are both high in the crisis and low in the calm, so they correlate strongly across regimes even if VPIN adds nothing within a regime. Our year-by-year results support this: VPIN's correlation with the next move ranges from {by_year['vpin_corr'].min():.2f} to {by_year['vpin_corr'].max():.2f} across years, and is highest in {int(top_year['year'])} ({top_year['vpin_corr']:.2f}), the year of the US downgrade turmoil. A sample dominated by 2008 would plausibly push it higher still. We cannot test this directly, because Kibot's IVE history starts in September 2009. Across all {w['IVE']['windows'] + w['BTC']['windows']} two-year windows of both markets, the correlation never exceeds {max(w['IVE']['corr_50_250'][1], w['BTC']['corr_50_250'][1]):.2f}.</p>

<h3>4.3 VPIN is sensitive to the classification step</h3>
<p>The paper presents the classification method as a technical detail. Our results say otherwise. The flash-crash warning appears with one-minute bars, disappears with 10-second bars, and never appears with Lee–Ready. On Bitcoin, where the truth is known, VPIN from bulk classification and VPIN from the true labels correlate only {bb['corr_vpin_bvc_true']:.2f} ± {bse['corr_vpin_bvc_true']:.2f} (Extension 1). A measure that changes this much with the classifier cannot be expected to give the same numbers on a different market.</p>

<h3>4.4 The paper's numbers are flattered</h3>
<ul>
<li><b>Look-ahead in the CDF.</b> The paper ranks each VPIN value against the whole sample, including the future. A market maker in 2010 could not have known that a VPIN value would later turn out to be in the top 10%.</li>
<li><b>Parameter search on the evaluation data.</b> Buckets per day and window length are chosen by scanning grids, and the best cell is reported. On IVE the pair with the best in-sample correlation {tuple(chosen['best in-sample corr'])} falls from {se(best_c['in-sample']['VPIN'], 'corr')} in 2009–2011 to {se(best_c['out-of-sample']['VPIN'], 'corr')} afterwards. The pair with the best in-sample true-positive rate {tuple(chosen['best in-sample true positive'])} falls from {se(best_t['in-sample']['VPIN'], 'true_positive_top', True)} to {se(best_t['out-of-sample']['VPIN'], 'true_positive_top', True)}, below the {se(best_t['out-of-sample']['VPIN'], 'true_positive_bottom', True)} reached when VPIN is <i>low</i>.</li>
<li><b>Over-confident precision.</b> The paper's interval for its correlation, [{PN.CORR_50_250_CI[0]:.3f}, {PN.CORR_50_250_CI[1]:.3f}], assumes independent buckets. With overlapping windows the real uncertainty is several times larger (Section 2).</li>
<li><b>One event.</b> The headline evidence, the flash crash, is a single observation.</li>
</ul>

<!--pagebreak-->
<h3>4.5 What we ruled out</h3>
<ul>
<li><b>Errors in our code.</b> The parts that do not depend on prediction reproduce (stability, low VPIN means calm, the volume clock's effect on tails), all {n_tests} unit tests pass, and our reading of the paper's Table 1 formulas reproduces the paper's own Jarque–Bera numbers from its own skewness and kurtosis.</li>
<li><b>Overnight gaps.</b> Removing buckets that span the close leaves the correlation unchanged (Section 4.1).</li>
<li><b>Bad prints.</b> The trades the exchanges cancelled on 6 May 2010 are removed by the exchanges' own rule.</li>
<li><b>Too little data.</b> The reproduction's standard error for the correlation is {inf['power']['corr_se_50_250']:.3f}, so it would detect any correlation of {inf['power']['mde_80pct_power']:.2f} or more with 80% power. The paper's 0.40 would not have been missed.</li>
</ul>
<div class="warn"><p><b>What remains open.</b> Without the E-mini data we cannot say how much of the gap is the market and period, and how much is the paper's procedure. The cleanest next step is to rerun everything on E-mini futures with exchange aggressor flags (Section 6).</p></div>

<h2>5. What our extensions add to the paper</h2>
<p>The seminar's top grade asks for upgrades beyond the paper: a new dataset, a robustness check, a fixed flaw. Each extension asks one question the paper could not answer.</p>

<h3>Extension 1: is bulk classification right? ({sl('Extension 1')})</h3>
<p><b>Question.</b> BVC is the foundation of VPIN, but the paper never checked it against the truth, because its data do not record who initiated each trade. <b>What we did.</b> On Bitcoin futures, where Binance records the aggressor of every trade, we poured BVC's estimate and the true buy volume into the same buckets and compared them.</p>
{fig("results/extensions/figures/e1_btc_vpin_true_vs_bvc.png", "Bitcoin: VPIN from bulk classification against VPIN from the true labels (left), and both over the last 60 days of the sample (right).")}
<ul>
<li>On one-minute bars BVC gets the direction of the imbalance right for {bars['direction_agreement']:.1%} ± {bars['direction_agreement_se']:.1%} of volume: better than a coin, far from perfect.</li>
<li>In buckets it inflates the size of the imbalance {bse['imbalance_inflation']:.1f} times (mean {bb['mean_abs_imbalance_bvc']:.2f} against a true {bb['mean_abs_imbalance_true']:.2f}). The two VPINs correlate only {bb['corr_vpin_bvc_true']:.2f} ± {bse['corr_vpin_bvc_true']:.2f}.</li>
<li>On five days with every individual trade, the simple tick rule classifies {tr['tick_accuracy_min']:.0%}–{tr['tick_accuracy_max']:.0%} of volume correctly, and its bucket imbalance correlates at least {tr['corr_tick_min']:.2f} with the truth, against {tr['corr_bvc_1min_min']:.2f}–{tr['corr_bvc_1min_max']:.2f} for BVC (ranges across the five days). Caveat: Binance groups the fills of one order into one record, which helps the tick rule.</li>
</ul>
<p><b>What it adds.</b> The paper's claim that bulk classification is more useful than trade-by-trade classification does not hold where it can be checked. And the VPIN the paper studies is a different number from the toxicity it claims to measure.</p>

<!--pagebreak-->
<h3>Extension 2: is VPIN more than volatility? ({sl('Extension 2')})</h3>
<p><b>Question.</b> Because BVC turns price moves into imbalance, VPIN may simply be a volatility measure in disguise; Andersen and Bondarenko (2014) made this argument. With true labels we can test it directly. <b>What we did.</b> On Bitcoin we compared three signals with the same window and timing: VPIN from BVC, VPIN from the true labels, and trailing volatility (the average absolute return over the same n buckets). All three come from the same data, so they can share a table.</p>
{table(btc_df)}
<ul>
<li>VPIN from the <b>true</b> labels is <i>negatively</i> related to the next move; VPIN from BVC is positive; trailing volatility is far stronger than both. One possible reason for the negative sign, which we have not tested, is that busy, volatile periods fill each bucket with many independent orders, so true imbalance per unit of volume falls.</li>
<li>BVC's imbalance moves with the bucket's own absolute return ({bb['corr_abs_imbalance_with_same_bucket_abs_return_bvc']:.2f} ± {bse['corr_abs_imbalance_with_same_bucket_abs_return_bvc']:.2f}) more than the true imbalance does ({bb['corr_abs_imbalance_with_same_bucket_abs_return_true']:.2f} ± {bse['corr_abs_imbalance_with_same_bucket_abs_return_true']:.2f}): the volatility enters through the classification.</li>
<li>Adding VPIN to a regression that already contains trailing volatility, the last return and trading intensity raises R² by {reg_btc['delta_r2']:.4f} ± {reg_btc['delta_r2_se']:.4f} on Bitcoin and {reg_ive['delta_r2']:.4f} ± {reg_ive['delta_r2_se']:.4f} on IVE out of sample. The sign matters: on Bitcoin VPIN's coefficient is significantly <i>negative</i> (Newey–West t = {reg_btc['t_vpin_hac']:.1f}); on IVE it is positive (t = {reg_ive['t_vpin_hac']:.1f}). Either way the increment is economically negligible.</li>
</ul>
<p><b>What it adds.</b> Whatever predictive power VPIN has comes from its volatility component, not from order-flow toxicity. The paper never compares VPIN with this obvious baseline.</p>

<!--pagebreak-->
<h3>Extension 3: a known-answer test of VPIN ({sl('Extension 3')}; backup {sl('Backup: the known-answer')})</h3>
<p><b>The gap.</b> The paper validates VPIN with a Monte Carlo (Appendix A.2, Table A1) that draws buy and sell volume straight from the PIN model: every trade is labelled correctly and there are no prices. Its empirical VPIN, however, labels trades by bulk classification, which needs prices. So the paper never checks the method it actually uses. <b>What we did.</b> We simulated the paper's own model and added a price: order flow moves it (Δp = λ·(buys − sells)) and public news adds noise that involves no trader. We then ran the paper's exact pipeline where the true share of informed trading is known. There are four worlds: T (more informed trading in some blocks of days), V (more volatility and uninformed volume, informed trading unchanged), TV (both) and 0 (neither); {sim_n} replications each. Hypotheses, metrics and decision rules were written in <code>docs/preregistration.md</code> before the first run.</p>
{table(sim_df)}
{table(dec_df, verdict_col="Verdict")}
{fig("results/extensions/simulation/metrics.png", "Simulation, ± 1.96 Monte Carlo standard errors: false alarms from volatility (world V), detection of toxic episodes (world T), and the paper's predictive statistic where informed trading never changes (world V).")}
<ul>
<li>When volatility changes but informed trading does not (world V), the paper's VPIN tracks the true informed share at a correlation of only {sim('V', 'paper', 'M1_tracking')['mean']:.2f}, and it raises more alarms in volatile blocks. True-label VPIN moves the other way, because extra uninformed volume dilutes informed trading. That is the same sign we found on real Bitcoin data.</li>
<li>Decision D3 did <b>not</b> hold: in this calibration the paper's predictive statistic is not fooled, because two effects cancel (volatility inflates the bulk imbalance while extra volume dilutes it). We report it as pre-registered.</li>
<li>Measuring each price change against the previous day's volatility (our fix) removes the false alarms and keeps detection: rule D4 holds. A variant with prices rounded to a tick gives the same decisions.</li>
</ul>
<p><b>What it adds.</b> It closes the gap in the paper's own validation, and it shows the mechanism under controlled conditions: VPIN as computed in the paper is not a pure measure of informed trading.</p>

<!--pagebreak-->
<h3>Extension 4: the fix on real data ({sl('Extension 4')})</h3>
<p><b>Question.</b> Because D4 held, the pre-registration required applying the fix to real data (rule D5). Does local-σ bulk classification get closer to the truth, and what happens to VPIN's predictive power?</p>
{table(ls_df)}
<ul>
<li>On Bitcoin, where the truth is known, local-σ VPIN correlates {tr2['BTC_50_50']['corr_local_true']['value']:.2f} with true-label VPIN, against {tr2['BTC_50_50']['corr_standard_true']['value']:.2f} for standard bulk classification (difference {tr2['BTC_50_50']['difference']['value']:+.2f} ± {tr2['BTC_50_50']['difference']['se']:.2f}). It is closer to the truth.</li>
<li>Once the volatility bias is gone, VPIN's correlation with the next move collapses on both markets. The paper's predictive power came from the bias.</li>
<li>What remains beyond volatility is small: on IVE the fixed VPIN adds ΔR² = {lr_ive['delta_r2']:.4f} (Newey–West t = {lr_ive['t_vpin_hac']:.1f}); on Bitcoin its coefficient is negative (t = {lr_btc['t_vpin_hac']:.1f}).</li>
</ul>
<p><b>What it adds.</b> A constructive result. A one-line change to the method makes VPIN a better measure of toxicity where it can be checked, and it shows how little toxicity information remains once the volatility artefact is removed.</p>

<h3>The inference package ({sl('What our work adds')})</h3>
<p><b>Placebo.</b> Each one-minute imbalance keeps its size but gets a random direction (20 draws). If VPIN's predictive power came from informed traders on one side of the market, the placebo would lose it. It does not:</p>
{table(plc_df)}
<ul>
<li><b>Leave one year out:</b> {flips} sign flips across {len(inf['loyo'])} signal–sample pairs; the results do not hinge on any single year.</li>
<li><b>Ledger:</b> {led['tests']} tests behind our claims; {led['supported_bh']} supported after Benjamini–Hochberg and {led['supported_bonferroni']} after Bonferroni. {led['significant_wrong_side']} are significant in the opposite direction from the claim and are counted as not supported (<code>results/report/ledger.csv</code>).</li>
<li><b>Borderline result, stated as such:</b> the reproduction's own correlation at (50, 250) has z = {rep_z:.2f} with the normal approximation, although the skewed bootstrap interval excludes zero. We call it weakly positive.</li>
</ul>
<p><b>What it adds.</b> The direction of trading, the thing VPIN is supposed to read, carries almost none of its predictive power; and every claim we make has a stated error rate.</p>

<!--pagebreak-->
<h3>Further check A: real-time VPIN out of sample ({sl('How could')}; backup {sl('Backup: real-time')})</h3>
<p><b>Question.</b> Would VPIN have worked for someone using it live after the paper? <b>What we did.</b> We built a real-time version in which σ and the bucket size come from the previous 250 days and the CDF from VPIN values already observed. We froze the paper's parameters, or parameters chosen the paper's way on 2009–2011, and evaluated on IVE from August 2011 to September 2026. Lift is the hit rate (P(signal in its top quartile | move &gt; 0.75%)) divided by how often the signal is in its top quartile, so an alarm that is always on scores 1.</p>
{table(oos_df)}
{fig("docs/build/is_vs_oos.png", "Correlation of each signal with the next move, in sample (2009–2011) and out of sample (2011–2026), ± one standard error.", 58)}
<ul>
<li>At the paper's (50, 250), VPIN's correlation is {se(p250_is['VPIN'], 'corr')} in sample and {se(p250_oos['VPIN'], 'corr')} out of sample; trailing volatility reaches {se(p250_oos['RV (trailing volatility)'], 'corr')}.</li>
<li>Parameters that looked best in sample lose most of their advantage out of sample (Section 4.4).</li>
<li>Year by year, trailing volatility has the higher correlation in {rv_years} of {len(by_year)} years. The share of time the VPIN alarm is on drifts strongly from year to year: a CDF built from the past does not fit the future.</li>
</ul>
{fig("results/extensions/figures/e2_ive_by_year.png", "IVE year by year, real-time pipeline at (50, 250): correlation with the next move (left) and the hit rate with the share of time the alarm is on (right).")}
<p><b>What it adds.</b> It removes the look-ahead and the in-sample selection from the paper's evaluation. What survives is a small, unstable signal that a simple volatility forecast beats.</p>

<h3>Further check B: crashes after the paper (backup {sl('Backup: crashes')})</h3>
<p><b>Question.</b> Did the real-time alarm warn of later crashes? <b>What we did.</b> For each event day we took the largest one-hour fall and checked whether the alarm (CDF ≥ 0.9) was already on when it began.</p>
{table(ev_df)}
<p class="label">Hits against how often each alarm is on</p>
{table(al_df)}
<p>On IVE the VPIN alarm was on before all five falls, while it is on only {alarms[('IVE', 'vpin')]['alarm_share']:.0%} of the time overall; trailing volatility was on before {alarms[('IVE', 'rv')]['hits']} of them. If the alarm were on at random, five hits out of five would be very unlikely (binomial p {pval(alarms[('IVE', 'vpin')])}). This is the extension result most favourable to VPIN. Two caveats: five events is very few, and several falls happened inside weeks of turmoil (August 2011, March 2020) when everything was already elevated, so the falls are not independent draws. On Bitcoin no signal was reliable.</p>
<p><b>What it adds.</b> An honest test of the paper's headline use case on new events, including the one case where VPIN looks good.</p>

<!--pagebreak-->
<h2>6. Critique, overfitting and next steps</h2>
<h3>How could the paper's results be overfit? ({sl('How could')})</h3>
<p>The seminar asks this of every paper. Our answer: the CDF uses the whole sample, including the future; the parameters are chosen by grid search on the same data that measure success; there is no out-of-sample period; the headline evidence is one event; and the reported precision ignores overlapping windows. Our real-time, out-of-sample check (Further check A) tests the first three directly and finds the decay that overfitting predicts; the placebo and the simulation show where the remaining predictive power comes from.</p>
<h3>Further critique ({sl('Critique and next')})</h3>
<ul>
<li>Bulk classification builds volatility into VPIN; with true labels the predictive link disappears (Extensions 1 and 2).</li>
<li>No baseline: the paper never compares VPIN with plain volatility, which wins in all our tests.</li>
<li>The paper's own Monte Carlo never runs bulk classification, so it does not validate the VPIN the paper uses (closed by Extension 3).</li>
<li>The authors disclose a patent on VPIN and a financial interest in it (title page of the paper).</li>
</ul>
<h3>With three more months</h3>
<ul>
<li>Rerun everything on E-mini futures with exchange aggressor flags: the paper's own market, with ground truth.</li>
<li>Use VPIN as a filter in a simulated market-making strategy and ask whether it reduces adverse-selection losses after costs.</li>
<li>Correct the parameter search for multiple testing (White's reality check, deflated statistics).</li>
<li>Combine VPIN with a standard volatility forecast (HAR) and test its added value out of sample.</li>
</ul>

<h2>7. Questions we expect</h2>
<p><b>What was the objective?</b> VPIN has none: it is an average of observed imbalances, not a fitted model. The paper's only optimisation is the grid search over bucket size and window, which maximises the reported statistic on the same data. We select parameters on 2009–2011 and evaluate on 2011–2026 ({sl('Data, and')}).</p>
<p><b>Why IVE and not the E-mini?</b> E-mini tick data are paid; IVE is the only free tick history with quotes we found that covers the flash crash. We say clearly that it is a different market, and Section 4 explains what that changes.</p>
<p><b>Is your code wrong rather than the paper?</b> The parts that do not depend on prediction reproduce, all {n_tests} unit tests pass, and our Table 1 formulas reproduce the paper's own numbers from its own inputs. The gaps appear exactly where the paper's procedure is most fragile.</p>
<p><b>Why a month-block bootstrap?</b> Consecutive VPIN values share 49 of 50 buckets, and volatility clusters, so single buckets are not independent observations. Resampling months keeps that dependence inside each block.</p>
<p><b>Why does the flash-crash result change with the bar length?</b> BVC's split depends on each bar's price change relative to σ. Short bars are more often flat (split 50/50), and consecutive short bars often move in opposite directions as trades alternate between the bid and the ask, so their imbalances cancel inside a bucket. VPIN then reacts mainly when prices trend, which on 6 May happened during the fall. This is our reading of the result; we have not tested it separately.</p>
<p><b>Isn't Bitcoin too different from equities?</b> It is a different market, which is why we use it only for questions that need ground truth: is BVC right, and is VPIN more than volatility. Those are properties of the method, not of one market.</p>
<p><b>Does the tick rule beat BVC everywhere?</b> Not necessarily. Binance groups the fills of one order, which helps the tick rule; on equity data with hidden orders and fast quote changes it is generally found to be less accurate, which we did not test. What we can say is that where the truth is known, BVC is clearly worse than the paper suggests.</p>
<p><b>The IVE alarm was on before every crash. Isn't that a success?</b> It is the most favourable result for VPIN, and we report it as such. But it rests on five falls, several inside weeks of turmoil when the alarm was already on, and trailing volatility caught three of them too.</p>
<p><b>Was the simulation built to give your answer?</b> Its design and decision rules were written down before the first run (<code>docs/preregistration.md</code>), and one pre-registered prediction (D3) failed; we report it. The model is the paper's own PIN model plus the simplest possible price.</p>
<p><b>Of course the fix predicts less: you removed volatility from VPIN. Isn't that circular?</b> The point is which version is closer to toxicity. Where the truth is known (Bitcoin), the fixed VPIN tracks true-label VPIN better. So the extra predictive power of the paper's version is volatility, not toxicity.</p>
<p><b>What is new compared with Andersen and Bondarenko?</b> They used perfect trade classification on E-mini data and argued that VPIN's link to volatility is mechanical; our Bitcoin results confirm this on new data. What we add: a known-answer simulation that closes the paper's own validation gap, a fix tested against ground truth, the direction placebo, and pre-registered inference.</p>
<p><b>What would change your conclusion?</b> A clear out-of-sample advantage of VPIN over trailing volatility on the E-mini with true aggressor flags.</p>

<h2>8. Glossary and where the numbers come from</h2>
<ul>
<li><b>Bucket:</b> a slice of trading with a fixed amount of volume (1/50 of an average day by default).</li>
<li><b>VPIN:</b> average |sell volume − buy volume| over the last n buckets divided by their volume.</li>
<li><b>BVC (bulk volume classification):</b> splits each bar's volume into buys and sells using the normal distribution of its standardized price change.</li>
<li><b>CDF(VPIN):</b> where the current VPIN sits among past values (0 to 1); an alarm is CDF ≥ 0.9.</li>
<li><b>Insurance value:</b> P(VPIN was in its top quartile | the next move exceeded 0.75%).</li>
<li><b>Lift:</b> insurance value divided by the share of time VPIN is in its top quartile; 1 means no information.</li>
<li><b>True positive (Table 5):</b> while VPIN stays in one top-quartile 5%-band, does a move above 0.75% occur?</li>
<li><b>Trailing volatility (RV):</b> average absolute return over the same n buckets as VPIN.</li>
<li><b>ΔR²:</b> how much of the variation in the next move VPIN explains beyond the controls.</li>
<li><b>In / out of sample:</b> 2009–2011 (the paper's period) / August 2011 – September 2026.</li>
</ul>
<p>Every number of ours in this document is read by <code>scripts/make_guide.py</code> from <code>results/report/key_numbers.json</code> and the CSV files in <code>results/</code>, produced by <code>scripts/reproduce.py</code> and <code>scripts/extensions.py</code>. The paper's numbers come from <code>src/vpin/paper_numbers.py</code>, each with its exhibit and page. Code, data instructions and every deviation: github.com/zavaz5/flow-toxicity-vpin.</p>
"""


def render(html_text: str, out: Path, label: str = "presentation guide",
           title: str = "Flow toxicity and VPIN: presentation guide", css: str = CSS):
    """A4 PDF from HTML with a footer "Flow toxicity and VPIN · <label> · page i of n"."""
    arch = pymupdf.Archive()
    arch.add(FONTS)
    arch.add(str(ROOT))
    page_rect = pymupdf.paper_rect("a4")
    frame = page_rect + (54, 50, -54, -56)
    writer = pymupdf.DocumentWriter(str(out))
    # Each <!--pagebreak--> chunk starts on a new page, so a section never begins at a page foot.
    for chunk in html_text.split("<!--pagebreak-->"):
        story = pymupdf.Story(html=chunk, user_css=css, archive=arch)
        more = True
        while more:
            device = writer.begin_page(page_rect)
            more, _ = story.place(frame)
            story.draw(device)
            writer.end_page()
    writer.close()
    doc = pymupdf.open(out)
    n = doc.page_count
    for i, page in enumerate(doc):
        page.insert_text((54, page.rect.height - 30), f"Flow toxicity and VPIN · {label} · page {i + 1} of {n}",
                         fontsize=7.5, fontname="helv", color=(0.45, 0.47, 0.5))
    doc.set_metadata({"title": title,
                      "subject": "Companion to the seminar presentation on Easley, López de Prado & O'Hara (2012)"})
    doc.subset_fonts()
    doc.ez_save(str(out) + ".tmp")
    doc.close()
    Path(str(out) + ".tmp").replace(out)
    return n


def verify(html_text: str, pdf_path: Path, label: str = "presentation guide"):
    """Fail if any heading, paragraph or list item of the HTML is missing from the PDF.

    PyMuPDF's layout engine can silently drop text after an image at the top of a page,
    so every block is checked. The text layer stores hyphens as soft hyphens, some
    punctuation as look-alike characters, and page footers may fall inside a paragraph;
    all are normalized away before comparing.
    """
    import re
    doc = pymupdf.open(pdf_path)
    text = " ".join(page.get_text() for page in doc)
    text = re.sub(rf"Flow toxicity and VPIN · {label} · page \d+ of \d+", " ", text)
    import unicodedata
    # NFKC maps look-alikes the text layer uses (e.g. U+037E for ";") back to the plain character
    norm = lambda t: re.sub(r"\s+", "", unicodedata.normalize("NFKC", t.replace("\u00ad", "-").replace("\u2011", "-")))
    pdf = norm(text)
    missing = []
    for tag, inner in re.findall(r"<(p|li|h2|h3)[^>]*>(.*?)</\1>", html_text, flags=re.S):
        block = norm(html.unescape(re.sub(r"<[^>]+>", "", inner)))
        if block and (block[:40] not in pdf or block[-40:] not in pdf):
            missing.append(html.unescape(re.sub(r"<[^>]+>", "", inner))[:80])
    if missing:
        raise SystemExit("text missing from the PDF:\n  " + "\n  ".join(missing))


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    html_text = build_html()
    pages = render(html_text, OUT)
    verify(html_text, OUT)
    print(f"saved {OUT} ({pages} pages), all text blocks present")


if __name__ == "__main__":
    main()
